import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def create_use_case_table(doc, table_num, table_title, uc_title, uc_goal, scenarios):
    # Table Title (e.g. Table 5 Submit PPMP)
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.line_spacing = 1.15
    r_t = p_title.add_run(f"Table {table_num} {table_title}")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(12)
    r_t.bold = True

    # 1-column table with 2 rows: Row 0 is Use Case Title header, Row 1 is Body
    table = doc.add_table(rows=2, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.0)

    # Set table borders: top single, bottom single, middle single, no left/right/vertical
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:left w:val="nil"/>\n'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:right w:val="nil"/>\n'
        f'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:insideV w:val="nil"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # Row 0: Use Case Header (e.g. Use Case 3: Submit PPMP)
    cell_0 = table.rows[0].cells[0]
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(4)
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(uc_title)
    r0.font.name = "Arial"
    r0.font.size = Pt(11)
    r0.bold = True

    # Row 1: Content
    cell_1 = table.rows[1].cells[0]
    
    # Goal paragraph
    p_goal = cell_1.paragraphs[0]
    p_goal.paragraph_format.space_before = Pt(4)
    p_goal.paragraph_format.space_after = Pt(2)
    p_goal.paragraph_format.line_spacing = 1.15
    rg1 = p_goal.add_run("Use Case Goal: ")
    rg1.bold = True
    rg1.font.name = "Arial"
    rg1.font.size = Pt(10.5)
    rg2 = p_goal.add_run(uc_goal)
    rg2.font.name = "Arial"
    rg2.font.size = Pt(10.5)

    for sc in scenarios:
        # Scenario Title
        p_sc = cell_1.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(4)
        p_sc.paragraph_format.space_after = Pt(1)
        p_sc.paragraph_format.line_spacing = 1.15
        r = p_sc.add_run(sc['name'])
        r.bold = True
        r.font.name = "Arial"
        r.font.size = Pt(10.5)

        # Actor
        p_act = cell_1.add_paragraph()
        p_act.paragraph_format.space_before = Pt(0)
        p_act.paragraph_format.space_after = Pt(1)
        p_act.paragraph_format.line_spacing = 1.15
        r_lbl = p_act.add_run("Actor: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_val = p_act.add_run(sc['actor'])
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)

        # Pre-condition
        p_pre = cell_1.add_paragraph()
        p_pre.paragraph_format.space_before = Pt(0)
        p_pre.paragraph_format.space_after = Pt(1)
        p_pre.paragraph_format.line_spacing = 1.15
        r_lbl = p_pre.add_run("Pre-condition: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_val = p_pre.add_run(sc['precondition'])
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)

        # Main Flow
        p_mf = cell_1.add_paragraph()
        p_mf.paragraph_format.space_before = Pt(0)
        p_mf.paragraph_format.space_after = Pt(1)
        p_mf.paragraph_format.line_spacing = 1.15
        r_lbl = p_mf.add_run("Main Flow:")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)

        # Steps
        for step in sc['steps']:
            p_st = cell_1.add_paragraph()
            p_st.paragraph_format.space_before = Pt(0)
            p_st.paragraph_format.space_after = Pt(1)
            p_st.paragraph_format.line_spacing = 1.15
            parts = step.split(':', 1)
            if len(parts) == 2:
                r_st_lbl = p_st.add_run(parts[0] + ':')
                r_st_lbl.bold = True
                r_st_lbl.font.name = "Arial"
                r_st_lbl.font.size = Pt(10.5)
                r_st_val = p_st.add_run(parts[1])
                r_st_val.font.name = "Arial"
                r_st_val.font.size = Pt(10.5)
            else:
                r_st = p_st.add_run(step)
                r_st.font.name = "Arial"
                r_st.font.size = Pt(10.5)

        # User Action & System Response pairs
        for action, response in sc.get('actions_responses', []):
            p_ua = cell_1.add_paragraph()
            p_ua.paragraph_format.space_before = Pt(0)
            p_ua.paragraph_format.space_after = Pt(1)
            p_ua.paragraph_format.line_spacing = 1.15
            r_ua_lbl = p_ua.add_run("User Action: ")
            r_ua_lbl.bold = True
            r_ua_lbl.font.name = "Arial"
            r_ua_lbl.font.size = Pt(10.5)
            r_ua_val = p_ua.add_run(action)
            r_ua_val.font.name = "Arial"
            r_ua_val.font.size = Pt(10.5)

            p_sr = cell_1.add_paragraph()
            p_sr.paragraph_format.space_before = Pt(0)
            p_sr.paragraph_format.space_after = Pt(2)
            p_sr.paragraph_format.line_spacing = 1.15
            r_sr_lbl = p_sr.add_run("System Response: ")
            r_sr_lbl.bold = True
            r_sr_lbl.font.name = "Arial"
            r_sr_lbl.font.size = Pt(10.5)
            r_sr_val = p_sr.add_run(response)
            r_sr_val.font.name = "Arial"
            r_sr_val.font.size = Pt(10.5)

    # Empty paragraph after table for spacing
    p_after = doc.add_paragraph()
    p_after.paragraph_format.space_before = Pt(0)
    p_after.paragraph_format.space_after = Pt(6)
    p_after.paragraph_format.line_spacing = 1.0

if __name__ == '__main__':
    doc_test = docx.Document()
    create_use_case_table(
        doc_test,
        table_num='5',
        table_title='Submit PPMP',
        uc_title='Use Case 3: Submit PPMP',
        uc_goal='To allow the Office User to submit a completed PPMP for executive review.',
        scenarios=[{
            'name': 'Scenario 1: Successful PPMP submission.',
            'actor': 'Office User',
            'precondition': 'The user is logged in, authorized as Office User, and the PPMP draft is complete and valid.',
            'steps': [
                'Step 1: The user opens a saved PPMP draft.',
                'Step 2: The system displays the complete PPMP summary.',
                'Step 3: The user reviews the entered information and clicks Submit PPMP.',
                'Step 4: The system validates completeness and required values.',
                'Step 5: The system updates the PPMP status to Submitted.',
                'Step 6: The system records the submission and notifies the Executive Officer.'
            ],
            'actions_responses': [
                ('Reviews the PPMP and clicks "Submit PPMP."', 'Validates the PPMP, updates its status to Submitted, and records the submission.')
            ]
        }]
    )
    doc_test.save('Documents/test_use_case_table.docx')
    print('Successfully generated test table matching exact format!')

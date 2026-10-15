import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

# Exact Use Cases matching the UML Use Case Diagram and the user's format:
EXACT_USE_CASES = [
    {
        "table_num": "3",
        "table_title": "Login",
        "uc_title": "Use Case 1: Login",
        "goal": "To allow authorized users (Farm Owner, Customer, Delivery Rider) to securely log in to the Aquaculture Value Chain Management System with Predictive Analytics.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Login.",
                "actor": "Farm Owner, Customer, Delivery Rider",
                "precondition": "The user is registered in the system with an active account and assigned role.",
                "steps": [
                    "Step 1: The user opens the system login page.",
                    "Step 2: The system displays the authentication interface with username and password fields.",
                    "Step 3: The user enters registered credentials and clicks \"Sign In\".",
                    "Step 4: The system validates credentials against the database.",
                    "Step 5: The system verifies the user role and generates an authenticated session.",
                    "Step 6: The system redirects the user to their role-specific dashboard."
                ],
                "user_action": "Enters registered credentials and clicks \"Sign In.\"",
                "system_response": "Validates the credentials, creates a secure session, and redirects the user to the role-based interface."
            },
            {
                "name": "Scenario 2: Failed Login Attempt.",
                "actor": "Farm Owner, Customer, Delivery Rider",
                "precondition": "The user enters invalid or unregistered credentials on the login page.",
                "steps": [
                    "Step 1: The user opens the system login page.",
                    "Step 2: The user enters incorrect username or password.",
                    "Step 3: The user clicks \"Sign In\".",
                    "Step 4: The system validates submitted credentials against the database.",
                    "Step 5: The system identifies mismatched credentials and logs the failed attempt.",
                    "Step 6: The system displays an authentication error message and prompts re-entry."
                ],
                "user_action": "Enters invalid credentials and clicks \"Sign In.\"",
                "system_response": "Detects invalid credentials, rejects authentication, and displays an error alert."
            }
        ]
    },
    {
        "table_num": "4",
        "table_title": "View Dashboard",
        "uc_title": "Use Case 2: View Dashboard",
        "goal": "To allow the Farm Owner to view real-time aquaculture operational summaries, pond telemetry, and pending orders.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Dashboard Overview.",
                "actor": "Farm Owner",
                "precondition": "The user is logged in and authorized as Farm Owner.",
                "steps": [
                    "Step 1: The user accesses the Executive Dashboard.",
                    "Step 2: The system queries active ponds, daily mortality, feed consumption, and order queues.",
                    "Step 3: The system renders summary KPI cards and live visual charts.",
                    "Step 4: The user reviews inventory alerts and pending dispatch requests.",
                    "Step 5: The user clicks on a KPI card to inspect detailed module logs.",
                    "Step 6: The system navigates to the selected management section."
                ],
                "user_action": "Accesses the Executive Dashboard and reviews operational indicators.",
                "system_response": "Retrieves real-time farm metrics and renders KPI overview cards and operational charts."
            }
        ]
    },
    {
        "table_num": "5",
        "table_title": "Create Stock Batch",
        "uc_title": "Use Case 3: Create Stock Batch",
        "goal": "To allow the Farm Owner to initialize and register a new aquaculture rearing batch (superworms or redclaw crayfish).",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Stock Batch Creation.",
                "actor": "Farm Owner",
                "precondition": "The user is logged in as Farm Owner, and rearing capacity is available.",
                "steps": [
                    "Step 1: The user opens the Batch Management module and clicks \"New Batch\".",
                    "Step 2: The system displays the batch creation form.",
                    "Step 3: The user enters species type, initial stocking count, stage, and start date.",
                    "Step 4: The system validates required input values and positive quantities.",
                    "Step 5: The system saves the new Cultivation_Batch record with a unique batch code.",
                    "Step 6: The system displays a confirmation message and updates batch rosters."
                ],
                "user_action": "Fills in batch specifications and clicks \"Save Batch.\"",
                "system_response": "Validates input values, creates the new batch record, and confirms successful initialization."
            }
        ]
    },
    {
        "table_num": "6",
        "table_title": "Assign Ponds",
        "uc_title": "Use Case 4: Assign Ponds",
        "goal": "To allow the Farm Owner to allocate juvenile batches to specific rearing ponds or cultivation shelves.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Pond Allocation.",
                "actor": "Farm Owner",
                "precondition": "The user is logged in as Farm Owner, and the target pond is sanitized and active.",
                "steps": [
                    "Step 1: The user opens the Pond Operations interface.",
                    "Step 2: The system displays the list of physical ponds and available capacities.",
                    "Step 3: The user selects a pond and assigns an active batch with density parameters.",
                    "Step 4: The system checks density limits and pond compatibility.",
                    "Step 5: The system updates the pond status to Occupied and links the batch record.",
                    "Step 6: The system records the assignment in the audit log and confirms the update."
                ],
                "user_action": "Selects target pond, assigns batch parameters, and clicks \"Assign Pond.\"",
                "system_response": "Validates pond capacity, links the batch to the pond, and updates the occupancy status."
            }
        ]
    },
    {
        "table_num": "7",
        "table_title": "Customer Purchase",
        "uc_title": "Use Case 5: Customer Purchase",
        "goal": "To allow Commercial Customers to browse the product catalog, add items to the cart, and select delivery or pickup options.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Catalog Browsing and Cart Checkout.",
                "actor": "Customer",
                "precondition": "The customer is on the public storefront, and commercial inventory is active.",
                "steps": [
                    "Step 1: The customer browses the product catalog (live crayfish, packaged superworms, feeds).",
                    "Step 2: The system displays available stock volumes and dual pricing options (per-piece / per-kg).",
                    "Step 3: The customer selects quantity, unit tier, and adds items to the shopping cart.",
                    "Step 4: The system calculates line-item subtotals and updates cart state.",
                    "Step 5: The customer selects fulfillment method (Delivery / Storefront Pickup) and enters address.",
                    "Step 6: The customer clicks \"Proceed to Checkout\" to submit the order request."
                ],
                "user_action": "Selects products, configures fulfillment method, and clicks \"Proceed to Checkout.\"",
                "system_response": "Calculates order total, reserves temporary stock, and forwards order details to payment processing."
            }
        ]
    },
    {
        "table_num": "8",
        "table_title": "Process Sales Order",
        "uc_title": "Use Case 6: Process Sales Order",
        "goal": "To allow the Farm Owner to review, approve, and finalize customer sales orders and verify payments.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Sales Order Approval.",
                "actor": "Farm Owner",
                "precondition": "A customer has placed an order with submitted payment proof (GCash receipt / COD).",
                "steps": [
                    "Step 1: The user opens the Orders Processing Queue.",
                    "Step 2: The system displays pending customer orders and uploaded payment attachments.",
                    "Step 3: The user reviews order items, delivery destination, and payment validity.",
                    "Step 4: The user clicks \"Approve Order\".",
                    "Step 5: The system updates Order status to Confirmed and Payment status to Approved.",
                    "Step 6: The system commits inventory deduction and queues the order for courier dispatch."
                ],
                "user_action": "Inspects order details and payment proof, then clicks \"Approve Order.\"",
                "system_response": "Validates payment verification, sets order status to Confirmed, and commits stock allocation."
            }
        ]
    },
    {
        "table_num": "9",
        "table_title": "Update Inventory",
        "uc_title": "Use Case 7: Update Inventory",
        "goal": "To allow the Farm Owner to adjust commercial product stock balances, configure dual pricing, and set reorder levels.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Stock Adjustment and Pricing Update.",
                "actor": "Farm Owner",
                "precondition": "The user is logged in as Farm Owner with inventory management permissions.",
                "steps": [
                    "Step 1: The user navigates to the Inventory Management screen.",
                    "Step 2: The system displays all commercial products, available stock, and unit prices.",
                    "Step 3: The user selects a product, enters adjusted stock counts or new per-piece/per-kg prices.",
                    "Step 4: The system validates entered numeric values.",
                    "Step 5: The system commits inventory modifications and creates a Stock_Movement entry.",
                    "Step 6: The system displays an update confirmation and refreshes the catalog pricing."
                ],
                "user_action": "Modifies inventory quantities or pricing parameters and clicks \"Save Changes.\"",
                "system_response": "Validates numeric inputs, updates product records, and logs the inventory movement."
            }
        ]
    },
    {
        "table_num": "10",
        "table_title": "Courier Dispatch",
        "uc_title": "Use Case 8: Courier Dispatch",
        "goal": "To allow the Farm Owner to assign confirmed orders to active Delivery Riders and allow riders to view delivery waypoints.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Courier Assignment.",
                "actor": "Farm Owner, Delivery Rider",
                "precondition": "An order is marked Confirmed, and active couriers are available.",
                "steps": [
                    "Step 1: The Farm Owner opens the Dispatch Pool and selects an unassigned confirmed order.",
                    "Step 2: The system displays active riders on duty.",
                    "Step 3: The Farm Owner selects a rider and clicks \"Assign Delivery\".",
                    "Step 4: The system creates a Delivery_Dispatch record and routes waypoints.",
                    "Step 5: The system notifies the assigned Delivery Rider with customer drop-off details.",
                    "Step 6: The Delivery Rider views the assigned delivery order on their mobile interface."
                ],
                "user_action": "Selects available courier and clicks \"Assign Delivery.\"",
                "system_response": "Creates dispatch assignment, updates delivery status to Dispatched, and alerts the assigned courier."
            }
        ]
    },
    {
        "table_num": "11",
        "table_title": "Confirm Delivery (Proof of Delivery)",
        "uc_title": "Use Case 9: Confirm Delivery (Proof of Delivery)",
        "goal": "To allow the Delivery Rider to capture recipient confirmation (POD photo/signature) and finalize order delivery.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Delivery Completion with POD.",
                "actor": "Delivery Rider",
                "precondition": "The order is in transit, and the rider has reached the customer drop-off location.",
                "steps": [
                    "Step 1: The Delivery Rider delivers physical packages and collects COD cash if applicable.",
                    "Step 2: The Delivery Rider opens the active order on their mobile interface.",
                    "Step 3: The Delivery Rider captures Proof of Delivery (recipient photo or signature).",
                    "Step 4: The Delivery Rider clicks \"Complete Delivery\".",
                    "Step 5: The system uploads POD evidence and updates Order status to Delivered.",
                    "Step 6: The system notifies the Customer and Farm Owner of successful fulfillment."
                ],
                "user_action": "Uploads POD photo/signature and clicks \"Complete Delivery.\"",
                "system_response": "Validates POD submission, updates order status to Delivered, and logs completion timestamp."
            }
        ]
    },
    {
        "table_num": "12",
        "table_title": "Export Report",
        "uc_title": "Use Case 10: Export Report",
        "goal": "To allow the Farm Owner to generate, preview, and print executive financial, inventory, and sales reports.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Report Generation and Export.",
                "actor": "Farm Owner",
                "precondition": "The user is logged in as Farm Owner, and operational records exist.",
                "steps": [
                    "Step 1: The user opens the Reports & Analytics module.",
                    "Step 2: The user selects report category (Sales Summary, Expense Breakdown, Inventory Log).",
                    "Step 3: The user defines date filters (daily, weekly, monthly) and clicks \"Generate Report\".",
                    "Step 4: The system queries database tables and compiles aggregate statistics.",
                    "Step 5: The system displays the formatted report preview.",
                    "Step 6: The user clicks \"Export PDF / Print\" to download official documentation."
                ],
                "user_action": "Selects report type, applies date range, and clicks \"Export PDF / Print.\"",
                "system_response": "Compiles transaction datasets, renders printable report layout, and triggers download."
            }
        ]
    },
    {
        "table_num": "13",
        "table_title": "Sales Forecasting",
        "uc_title": "Use Case 11: Sales Forecasting",
        "goal": "To allow the Farm Owner to execute predictive machine learning models for short-term revenue forecasting and customer segmentation.",
        "scenarios": [
            {
                "name": "Scenario 1: Successful Predictive Analytics Computation.",
                "actor": "Farm Owner",
                "precondition": "Historical sales and harvesting records are present in the database.",
                "steps": [
                    "Step 1: The user opens the Predictive Analytics Dashboard.",
                    "Step 2: The system loads historical sales records and pond harvest yields.",
                    "Step 3: The user clicks \"Run Predictive Models\".",
                    "Step 4: The system executes Linear Regression algorithms to forecast future sales trends.",
                    "Step 5: The system runs K-Means clustering to segment customer purchasing cohorts.",
                    "Step 6: The system visualizes projected revenue graphs and regional demand heatmaps."
                ],
                "user_action": "Navigates to Predictive Analytics and clicks \"Run Predictive Models.\"",
                "system_response": "Processes historical datasets through regression and clustering models, rendering forecast visualizations."
            }
        ]
    }
]

def render_use_case_table(doc, uc):
    # Table Title: Table X [Title]
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.line_spacing = 1.15
    r_t = p_title.add_run(f"Table {uc['table_num']} {uc['table_title']}")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(12)
    r_t.bold = True

    # 1-Column Table (2 rows: Row 0 is Use Case Title, Row 1 is Body)
    table = doc.add_table(rows=2, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.0)

    # Borders: top, inside horizontal, bottom. No left/right/vertical.
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:left w:val="nil"/>\n'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:right w:val="nil"/>\n'
        f'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:insideV w:val="nil"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # Row 0: Use Case Header
    cell_0 = table.rows[0].cells[0]
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(4)
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(uc['uc_title'])
    r0.font.name = "Arial"
    r0.font.size = Pt(11)
    r0.bold = True

    # Row 1: Content
    cell_1 = table.rows[1].cells[0]

    # Goal
    p_goal = cell_1.paragraphs[0]
    p_goal.paragraph_format.space_before = Pt(4)
    p_goal.paragraph_format.space_after = Pt(2)
    p_goal.paragraph_format.line_spacing = 1.15
    rg1 = p_goal.add_run("Use Case Goal: ")
    rg1.bold = True
    rg1.font.name = "Arial"
    rg1.font.size = Pt(10.5)
    rg2 = p_goal.add_run(uc['goal'])
    rg2.font.name = "Arial"
    rg2.font.size = Pt(10.5)

    for sc_idx, sc in enumerate(uc['scenarios']):
        # Scenario Title
        p_sc = cell_1.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(4 if sc_idx > 0 else 2)
        p_sc.paragraph_format.space_after = Pt(1)
        p_sc.paragraph_format.line_spacing = 1.15
        r_sc = p_sc.add_run(sc['name'])
        r_sc.bold = True
        r_sc.font.name = "Arial"
        r_sc.font.size = Pt(10.5)

        # Actor
        p_act = cell_1.add_paragraph()
        p_act.paragraph_format.space_before = Pt(0)
        p_act.paragraph_format.space_after = Pt(1)
        p_act.paragraph_format.line_spacing = 1.15
        r_lbl = p_act.add_run("Actor: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_val = p_act.add_run(sc['actor'])
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)

        # Pre-condition
        p_pre = cell_1.add_paragraph()
        p_pre.paragraph_format.space_before = Pt(0)
        p_pre.paragraph_format.space_after = Pt(1)
        p_pre.paragraph_format.line_spacing = 1.15
        r_lbl = p_pre.add_run("Pre-condition: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)
        r_val = p_pre.add_run(sc['precondition'])
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10.5)

        # Main Flow
        p_mf = cell_1.add_paragraph()
        p_mf.paragraph_format.space_before = Pt(0)
        p_mf.paragraph_format.space_after = Pt(1)
        p_mf.paragraph_format.line_spacing = 1.15
        r_lbl = p_mf.add_run("Main Flow:")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10.5)

        # Steps
        for step in sc['steps']:
            p_st = cell_1.add_paragraph()
            p_st.paragraph_format.space_before = Pt(0)
            p_st.paragraph_format.space_after = Pt(1)
            p_st.paragraph_format.line_spacing = 1.15
            parts = step.split(':', 1)
            if len(parts) == 2:
                r_lbl = p_st.add_run(parts[0] + ':')
                r_lbl.bold = True
                r_lbl.font.name = "Arial"
                r_lbl.font.size = Pt(10.5)
                r_val = p_st.add_run(parts[1])
                r_val.font.name = "Arial"
                r_val.font.size = Pt(10.5)
            else:
                r_st = p_st.add_run(step)
                r_st.font.name = "Arial"
                r_st.font.size = Pt(10.5)

        # User Action
        p_ua = cell_1.add_paragraph()
        p_ua.paragraph_format.space_before = Pt(0)
        p_ua.paragraph_format.space_after = Pt(1)
        p_ua.paragraph_format.line_spacing = 1.15
        r_ua_lbl = p_ua.add_run("User Action: ")
        r_ua_lbl.bold = True
        r_ua_lbl.font.name = "Arial"
        r_ua_lbl.font.size = Pt(10.5)
        r_ua_val = p_ua.add_run(sc['user_action'])
        r_ua_val.font.name = "Arial"
        r_ua_val.font.size = Pt(10.5)

        # System Response
        p_sr = cell_1.add_paragraph()
        p_sr.paragraph_format.space_before = Pt(0)
        p_sr.paragraph_format.space_after = Pt(2)
        p_sr.paragraph_format.line_spacing = 1.15
        r_sr_lbl = p_sr.add_run("System Response: ")
        r_sr_lbl.bold = True
        r_sr_lbl.font.name = "Arial"
        r_sr_lbl.font.size = Pt(10.5)
        r_sr_val = p_sr.add_run(sc['system_response'])
        r_sr_val.font.name = "Arial"
        r_sr_val.font.size = Pt(10.5)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
    p_sp.paragraph_format.line_spacing = 1.0

if __name__ == '__main__':
    doc_out = docx.Document()
    for section in doc_out.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    p_h = doc_out.add_paragraph()
    p_h.paragraph_format.space_before = Pt(0)
    p_h.paragraph_format.space_after = Pt(6)
    p_h.paragraph_format.line_spacing = 1.5
    r_h = p_h.add_run("3.6 Use Case Scenario")
    r_h.font.name = "Arial"
    r_h.font.size = Pt(12)
    r_h.bold = True

    p_i = doc_out.add_paragraph()
    p_i.paragraph_format.space_before = Pt(0)
    p_i.paragraph_format.space_after = Pt(6)
    p_i.paragraph_format.line_spacing = 1.5
    p_i.paragraph_format.first_line_indent = Inches(0.5)
    r_i = p_i.add_run(
        "The Use Case Scenario describes the detailed interactions between the primary actors (Farm Owner, "
        "Delivery Rider, and Commercial Customers) and the Aquaculture Value Chain Management System with Predictive Analytics. "
        "Each table corresponds to a major use case modeled in the UML Use Case Diagram, detailing the use case goal, "
        "operational pre-conditions, main step-by-step workflow, and summary user action and system response."
    )
    r_i.font.name = "Arial"
    r_i.font.size = Pt(12)

    for uc in EXACT_USE_CASES:
        render_use_case_table(doc_out, uc)

    doc_out.save("Documents/Use_Case_Scenario.docx")
    print(f"Successfully generated Documents/Use_Case_Scenario.docx with {len(EXACT_USE_CASES)} exact Use Case Tables!")

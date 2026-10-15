import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

def generate_identical_chapter1():
    doc = docx.Document()

    # Short bond paper (Letter: 8.5" x 11.0")
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    # Base Normal Style: Arial 12pt
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = None
    normal_style.paragraph_format.space_before = Pt(10)
    normal_style.paragraph_format.space_after = Pt(10)

    def p_title_line(text, bold=True, size=12):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = None
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Arial'
        run.font.size = Pt(size)
        return p

    def p_blank(count=1):
        for _ in range(count):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = None

    def p_chapter_head(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = None
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = None
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_body(text, indent=True):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.5)
        else:
            p.paragraph_format.first_line_indent = Inches(0.0)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = None
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_definition_item(bold_term, definition_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = None
        
        run_bold = p.add_run(bold_term)
        run_bold.bold = True
        run_bold.font.name = 'Arial'
        run_bold.font.size = Pt(12)
        
        run_text = p.add_run(definition_text)
        run_text.font.name = 'Arial'
        run_text.font.size = Pt(12)
        return p

    def p_significance_item(bold_stakeholder, description_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(10)
        p.paragraph_format.line_spacing = None
        
        run_bold = p.add_run(bold_stakeholder)
        run_bold.bold = True
        run_bold.font.name = 'Arial'
        run_bold.font.size = Pt(12)
        
        run_text = p.add_run(description_text)
        run_text.font.name = 'Arial'
        run_text.font.size = Pt(12)
        return p

    def p_objective(num_prefix, text, sub=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = None
        if sub:
            p.paragraph_format.left_indent = Inches(0.75)
            p.paragraph_format.first_line_indent = Inches(-0.35)
        else:
            p.paragraph_format.left_indent = Inches(0.4)
            p.paragraph_format.first_line_indent = Inches(-0.25)
        
        run_num = p.add_run(num_prefix)
        run_num.font.name = 'Arial'
        run_num.font.size = Pt(12)
        
        run_text = p.add_run(text)
        run_text.font.name = 'Arial'
        run_text.font.size = Pt(12)
        return p

    # ==================== TITLE PAGE ====================
    p_title_line("AQUACULTURE VALUE CHAIN MANAGEMENT SYSTEM WITH PREDICTIVE ANALYTICS", bold=True, size=12)
    p_blank(3)
    p_title_line("A Capstone Project", bold=False)
    p_title_line("Presented to the Faculty of the", bold=False)
    p_title_line("College of Computer Studies", bold=False)
    p_title_line("Carlos Hilado Memorial State University", bold=False)
    p_blank(3)
    p_title_line("In Partial Fulfilment", bold=False)
    p_title_line("of the Requirements for the Degree", bold=False)
    p_title_line("BACHELOR OF SCIENCE IN INFORMATION SYSTEMS", bold=True)
    p_blank(3)
    p_title_line("By", bold=False)
    p_blank(1)
    p_title_line("RESEARCH PROPONENTS", bold=True)
    p_blank(3)
    p_title_line("OCTOBER 2026", bold=True)

    # Page break to Chapter 1
    doc.add_page_break()

    # ==================== CHAPTER I ====================
    p_chapter_head("CHAPTER I")
    p_chapter_head("INTRODUCTION")

    p_body("Superworm and crayfish production play an emerging role in sustainable food production, animal feed development, and livelihood generation, particularly in small-scale and rural farming communities. Superworms (Zophobas morio) are increasingly recognized as a promising insect-based protein source for animal feed and other nutritional applications, while crayfish, particularly freshwater species such as redclaw crayfish (Cherax quadricarinatus), are cultivated as high-value freshwater crustaceans in aquaculture systems (Rumbos & Athanassiou, 2021; Rigg et al., 2020). As demand for alternative protein sources and freshwater aquatic products continues to rise, the expansion of superworm and crayfish farming has highlighted the need for more efficient management across breeding, feeding, harvesting, inventory, and distribution systems. Digital tools are becoming increasingly important in improving productivity, resource monitoring, and decision-making in aquatic food and farm production systems (Food and Agriculture Organization, 2022; WorldFish, 2021). Aquaculture value chain research also emphasizes the importance of understanding the linkages among producers, input suppliers, market intermediaries, and consumers across production, processing, and distribution activities (Bush et al., 2019).", indent=True)

    p_body("One of the major challenges faced by superworm and crayfish farms is the lack of centralized systems capable of managing operational data across different production units, species, batches, and growth cycles. Without an integrated digital platform, farm managers may encounter difficulties in consolidating production records, monitoring feed and water usage, tracking mortality and growth performance, and evaluating overall farm productivity. Many small and medium farm enterprises continue to experience inefficient data management and limited technological integration. Digital platforms can support farmers and value chain actors by improving access to production information, feed management, market linkages, and farm performance tracking (WorldFish, 2021; WorldFish, 2022).", indent=True)

    p_body("To address these challenges, the development of a Superworm and Crayfish Value Chain Management System with Predictive Analytics is proposed to support more efficient data management and value chain coordination. The proposed system aims to provide a centralized digital platform capable of storing, organizing, and analyzing data related to superworm rearing, crayfish pond operations, feeding activities, inventory, harvesting, supply transactions, and market distribution. By incorporating predictive analytics, the system can generate data-driven insights that support planning and strategic decision-making, such as estimating production trends, forecasting harvest volume, monitoring supply levels, and improving resource allocation. Recent research in aquatic production forecasting demonstrates that predictive models using historical and environmental data can support producers and decision-makers in improving planning and resource management across aquaculture-related value chains (Hu et al., 2025).", indent=True)

    # 1.1 Background of the Study
    p_heading_3("1.1 Background of the Study")
    p_body("Silay Superworm & Crayfish, a farming enterprise located in Brgy. Lantad, Silay City, operates production areas for superworm rearing and crayfish cultivation. As the demand for alternative protein sources and freshwater aquaculture products continues to grow, the farm needs to modernize its operations to meet increasing market demand and manage growing production complexity. Stakeholders, including the farm owner, delivery riders, and customers, require timely and accurate information to monitor production, manage resources, and coordinate activities across different production and distribution channels.", indent=True)
    p_body("Currently, Silay Superworm & Crayfish faces challenges in managing operational data and coordinating value chain activities due to fragmented record-keeping practices. These include manual logs and paper-based systems, which limit efficient monitoring of superworm batches, crayfish pond performance, inventory tracking, feeding schedules, harvesting activities, and sales records. This fragmented approach reduces transparency and slows down decision-making, potentially leading to inefficient resource allocation and reduced competitiveness.", indent=True)
    p_body("To address these issues, a centralized digital management system with predictive analytics is proposed. The system will integrate superworm production management, crayfish pond monitoring, feed and harvest data, inventory records, and sales information into a unified web-based platform. By providing real-time insights and predictive forecasts, the system aims to improve coordination, enhance productivity, support better decision-making, and contribute to the sustainable growth of Silay Superworm & Crayfish.", indent=True)

    # 1.2 Objectives of the Study
    p_heading_3("1.2 Objectives of the Study")
    p_body("The purpose of the study is to develop a Silay Superworm & Crayfish Value Chain Management System with Predictive Analytics that improves the monitoring, management, and analysis of farm operations, including production, inventory, sales, and expenses.", indent=True)
    p_body("Specifically, this study aims to:", indent=True)
    
    p_objective("1. ", "Design and develop a Silay Superworm & Crayfish Value Chain Management System with the following features:", sub=False)
    p_objective("1.1. ", "Documentation of superworm and crayfish production processes, including breeding, feeding, harvesting, processing, distribution, and delivery;", sub=True)
    p_objective("1.2. ", "Sales and inventory management for tracking stock, transactions, and product availability;", sub=True)
    p_objective("1.3. ", "Product classification based on type, category, and scientific name;", sub=True)
    p_objective("1.4. ", "Monitoring of monthly operational expenses;", sub=True)
    p_objective("1.5. ", "Forecasting and predictive analytics for production, expenses, and market demand;", sub=True)
    p_objective("1.6. ", "Customer transaction management, including wholesale, retail, and reservations;", sub=True)
    p_objective("1.7. ", "Data clustering for analyzing production, sales patterns, and customer behavior; and", sub=True)
    p_objective("1.8. ", "Report generation and analytics for decision-making.", sub=True)
    p_objective("2. ", "Test the functionality of the developed system.", sub=False)
    p_objective("3. ", "Evaluate the usability of the system in terms of:", sub=False)
    p_objective("3.1. ", "System usefulness;", sub=True)
    p_objective("3.2. ", "Quality of information;", sub=True)
    p_objective("3.3. ", "Interface quality; and", sub=True)
    p_objective("3.4. ", "Overall usability.", sub=True)
    p_objective("4. ", "Develop a User Manual for the system.", sub=False)

    # 1.3 Scope and Limitations of the Study
    p_heading_3("1.3 Scope and Limitations of the Study")
    p_body("This study focuses on the development of a Silay Superworm & Crayfish Value Chain Management System with Predictive Analytics designed for Silay Superworm & Crayfish. The system provides a centralized platform for managing farm production data, including superworm rearing, crayfish pond operations, harvest records, product classification, inventory, sales transactions, and operational expenses. It also includes predictive analytics to generate forecasts related to production, expenses, and market demand. The system supports customer transaction management and generates analytical reports for decision-making. The system will be web-based and accessible through standard browsers, with data stored in a centralized database.", indent=True)
    p_body("The study has limitations. The system will not include real-time sensor integration, such as temperature, dissolved oxygen, pH levels, or humidity monitoring, since data will be manually entered by authorized users. Predictive analytics will rely on historical data and will not include advanced machine learning models. Furthermore, the system is tailored specifically for Silay Superworm & Crayfish and may require modifications if applied to other farming or aquaculture enterprises.", indent=True)

    # 1.4 Significance of the Study
    p_heading_3("1.4 Significance of the Study")
    p_body("The development of the Silay Superworm & Crayfish Value Chain Management System with Predictive Analytics is expected to benefit several stakeholders involved in superworm and crayfish production and distribution operations.", indent=True)
    p_significance_item("Farm Owners.", " The system provides a centralized executive management platform that unifies production monitoring, inventory balances, sales turnover, and operational expenses into comprehensive reports. It assists the farm owner in making timely and data-driven business decisions through predictive yield estimations, expense tracking, and market demand forecasts.")
    p_significance_item("Delivery Riders.", " The system provides a dedicated mobile delivery interface that streamlines order dispatching, provides accurate customer delivery addresses and contact information, facilitates navigation, and allows the capture of photo-verified Proof of Delivery upon successful product turnover.")
    p_significance_item("Customers and Buyers.", " The system offers an accessible and transparent online marketplace where customers can browse live superworm and crayfish inventories, select flexible fulfillment options (such as pickup or home delivery with GPS location pinning), and complete secure transactions through digital payment gateways or cash on delivery.")
    p_significance_item("Researchers and Future Developers.", " This study provides an empirical reference and architectural framework for future research on the digital transformation of small-scale agricultural enterprises, aquaculture supply chain automation, and the integration of predictive analytics in specialized farm management systems.")
    p_significance_item("The University and Academic Community.", " The study contributes to the academic literature in information systems by demonstrating the practical application of web development technologies, data analytics, and value chain modernization to enhance productivity and sustainability in local agro-aquaculture businesses.")

    # 1.5 Definition of Terms
    p_heading_3("1.5 Definition of Terms")
    p_definition_item("Aquaculture.", " The controlled farming, cultivation, and breeding of aquatic organisms in managed water environments. In this study, it specifically refers to the freshwater pond and tank rearing of Redclaw Crayfish (Cherax quadricarinatus) and the coordinated rearing of Superworms (Zophobas morio) at Silay Superworm & Crayfish in Brgy. Lantad, Silay City.")
    p_definition_item("Aquaculture Management System.", " An integrated web-based software platform engineered to centralize and automate farm operations. It integrates production batch records, pond water quality logs, feeding schedules, stock inventory, customer sales transactions, and operational expense tracking into a unified interface.")
    p_definition_item("Clustering.", " A data analysis and machine learning technique used to partition operational and sales records into distinct groups based on shared characteristics. In the system, clustering is used to categorize production batches, sales velocity, and customer purchasing patterns to identify top-performing products and key customer segments.")
    p_definition_item("Inventory Management.", " The systematic process of monitoring, controlling, and updating the quantities and movement of products within the farm. In this study, it refers to the digital ledger that tracks live stock balances of superworm boxes, crayfish juveniles, and breeder stocks, automatically updating quantities as batches are harvested or orders are fulfilled.")
    p_definition_item("Operational Expenses.", " The recurring expenditures necessary for maintaining the daily operations and production cycles of the farm. In this study, it encompasses recorded costs for commercial feeds, electricity, water utilities, labor wages, pond sanitization supplies, and delivery logistics.")
    p_definition_item("Predictive Analytics.", " A data analysis technique that evaluates historical production, expenditure, and sales data to recognize patterns and generate forecasts about future outcomes. In this study, it refers to the computational models used to estimate upcoming harvest yields, projected monthly operating costs, and forward product demand.")
    p_definition_item("Retail Sales.", " The commercial selling of superworms, crayfish, and related products in smaller quantities directly to individual consumers, hobbyists, anglers, and pet owners through the system's online marketplace.")
    p_definition_item("Value Chain.", " The comprehensive series of interrelated activities involved in producing and delivering farm goods to the final consumer. In this study, it encompasses breeding, feeding, growth monitoring, harvesting, processing, inventory warehousing, marketing, sales processing, and delivery distribution.")
    p_definition_item("Wholesale Sales.", " The commercial sale and bulk distribution of superworms and crayfish in large quantities to commercial retailers, pet shops, fish farms, and agricultural distributors, typically governed by volume-tiered pricing and scheduled logistics.")

    output_path = r'c:\Users\kymer\Desktop\Lantad ACarwfish Django\Documents\Final_Paper\Chapter_1.docx'
    doc.save(output_path)
    print("Successfully built exact match Chapter_1.docx with 200 twips (10pt) space before/after and auto line spacing!")

if __name__ == '__main__':
    generate_identical_chapter1()

import docx

doc = docx.Document(r'c:\Users\kymer\Desktop\Lantad ACarwfish Django\Documents\Final_Paper\Chapter_1.docx')

for i in range(24, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    pf = p.paragraph_format
    runs = [(r.text[:25], r.bold) for r in p.runs]
    print(f"P{i:02d} (align={p.alignment}, indent={pf.first_line_indent.inches if pf.first_line_indent else 0}in): '{p.text[:60]}...' | runs={runs}")

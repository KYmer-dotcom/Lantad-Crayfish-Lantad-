import docx

doc = docx.Document(r'c:\Users\kymer\Desktop\Lantad ACarwfish Django\Documents\Final_Paper\Chapter_1.docx')

for i in range(len(doc.paragraphs)-4, len(doc.paragraphs)):
    p = doc.paragraphs[i]
    print(f"=== P{i}: {p.text[:30]} ===")
    print(p._p.pPr.xml)

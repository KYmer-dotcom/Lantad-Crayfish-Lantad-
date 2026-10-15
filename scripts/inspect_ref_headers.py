import docx

doc = docx.Document(r'c:\Users\kymer\Desktop\Lantad ACarwfish Django\Documents\Final_Paper\Herbias, Hulleza, Opjer - GovSupply Connect_ A Municipal Procurement Transparency and Supply Distribution Monitoring System.docx')

print("Total sections:", len(doc.sections))
for i, s in enumerate(doc.sections):
    h = s.header
    f = s.footer
    print(f"Section {i}: Header paragraphs={len(h.paragraphs)}, text='{[p.text for p in h.paragraphs]}'")
    print(f"Section {i}: Footer paragraphs={len(f.paragraphs)}, text='{[p.text for p in f.paragraphs]}'")

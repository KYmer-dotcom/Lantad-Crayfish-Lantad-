import docx

doc = docx.Document(r'c:\Users\kymer\Desktop\Lantad ACarwfish Django\Documents\Final_Paper\Herbias, Hulleza, Opjer - GovSupply Connect_ A Municipal Procurement Transparency and Supply Distribution Monitoring System.docx')

print("Total paragraphs:", len(doc.paragraphs))
for i in range(min(45, len(doc.paragraphs))):
    p = doc.paragraphs[i]
    runs = [(r.text, r.font.name, r.font.size.pt if r.font.size else None, r.bold, r.italic) for r in p.runs]
    print(f"P{i:02d} [Style={p.style.name}, Align={p.alignment}]: '{p.text}' | Runs: {runs}")

# Check section margins and default styles
sec = doc.sections[0]
print(f"Margins: top={sec.top_margin.inches}, bottom={sec.bottom_margin.inches}, left={sec.left_margin.inches}, right={sec.right_margin.inches}")

norm = doc.styles['Normal']
print(f"Normal style font: {norm.font.name}, size: {norm.font.size.pt if norm.font.size else None}")

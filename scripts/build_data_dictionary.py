import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_apa_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="none"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_row_bottom_border(row):
    for cell in row.cells:
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

def add_data_table(doc, table_num, table_name, rows_data):
    # Title above table
    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(14)
    p_num.paragraph_format.space_after = Pt(2)
    p_num.paragraph_format.line_spacing = 1.0
    r_num = p_num.add_run(f"Table {table_num}")
    r_num.font.name = "Arial"
    r_num.font.size = Pt(11)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.paragraph_format.line_spacing = 1.0
    r_title = p_title.add_run(table_name)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(11)
    r_title.font.bold = False
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    # Table with 6 columns
    col_widths = [Inches(1.2), Inches(0.9), Inches(0.85), Inches(0.65), Inches(1.8), Inches(1.1)]
    table = doc.add_table(rows=len(rows_data) + 1, cols=6)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_apa_table_borders(table)

    headers = ["Field Name", "Data Type", "Data Format", "Field Size", "Description", "Example"]
    
    # Header Row
    hdr_row = table.rows[0]
    set_row_bottom_border(hdr_row)
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(heading)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Data Rows
    for row_idx, data in enumerate(rows_data):
        row = table.rows[row_idx + 1]
        for col_idx, text_val in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(text_val))
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0, 0, 0)
            if col_idx == 0 and ("(PK)" in str(text_val) or "(FK)" in str(text_val)):
                r.font.bold = False

    # Spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(8)
    p_sp.paragraph_format.line_spacing = 1.0

# 15 Core Entity Data Tables
DATA_DICTIONARY_TABLES = [
    {
        "num": 16,
        "name": "Users Data Table (AUTH_USER)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique auto-incremented identifier for user accounts.", "1"),
            ("username", "VARCHAR", "Alphanumeric", "150", "Unique system login credential handle.", "owner_silay"),
            ("password", "VARCHAR", "Encrypted Text", "128", "Argon2/PBKDF2 salted hash password string.", "pbkdf2_sha256$..."),
            ("email", "VARCHAR", "Email", "254", "Primary contact and notification email address.", "owner@lantadfarm.ph"),
            ("first_name", "VARCHAR", "Text", "150", "Given legal first name of the user.", "Eduardo"),
            ("last_name", "VARCHAR", "Text", "150", "Legal family surname of the user.", "Montinola"),
            ("role", "VARCHAR", "Text", "20", "Authorization level (OWNER, CUSTOMER, RIDER).", "OWNER"),
            ("is_active", "BOOLEAN", "Logical Flag", "1", "Flag specifying if account can log in and transact.", "1"),
            ("date_joined", "DATETIME", "Date & time", "—", "System timestamp of original account registration.", "2025-01-10 08:30:00"),
        ]
    },
    {
        "num": 17,
        "name": "Customers Data Table (SALES_CUSTOMER)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique auto-incremented customer profile identifier.", "101"),
            ("user_id (FK)", "INT", "Number", "11", "One-to-one foreign key referencing AUTH_USER(id).", "4"),
            ("name", "VARCHAR", "Text", "200", "Full individual name or commercial business name.", "Bacolod Seafood Grill"),
            ("customer_type", "VARCHAR", "Text", "20", "Client category (retailer, wholesaler, restaurant).", "restaurant"),
            ("phone", "VARCHAR", "Text", "20", "Active mobile contact number for SMS and delivery.", "09171234567"),
            ("address", "TEXT", "Long text", "—", "Complete delivery street address and landmark.", "Lacson St., Bacolod City"),
            ("map_latitude", "DECIMAL(9,6)", "Coordinate", "9,6", "GPS latitude coordinate for GIS route mapping.", "10.676500"),
            ("map_longitude", "DECIMAL(9,6)", "Coordinate", "9,6", "GPS longitude coordinate for GIS route mapping.", "122.950900"),
            ("credit_limit", "DECIMAL(12,2)", "Currency", "12,2", "Maximum allowable accounts receivable credit cap.", "25000.00"),
            ("created_at", "DATETIME", "Date & time", "—", "Timestamp when customer record was created.", "2025-01-15 09:00:00"),
        ]
    },
    {
        "num": 18,
        "name": "Delivery Riders Data Table (SALES_RIDER)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique auto-incremented courier identifier.", "1"),
            ("user_id (FK)", "INT", "Number", "11", "One-to-one foreign key referencing AUTH_USER(id).", "5"),
            ("name", "VARCHAR", "Text", "150", "Full legal name of the delivery personnel.", "Ramon Dela Cruz"),
            ("phone", "VARCHAR", "Text", "20", "Mobile phone number of the rider for coordination.", "09287654321"),
            ("vehicle_type", "VARCHAR", "Text", "50", "Dispatched transport mode (Motorcycle, Multicab).", "Motorcycle with Insulated Box"),
            ("plate_number", "VARCHAR", "Alphanumeric", "50", "Official LTO vehicle license plate number.", "NEG-4582"),
            ("status", "VARCHAR", "Text", "20", "Current operational state (AVAILABLE, BUSY, OFFLINE).", "AVAILABLE"),
            ("is_active", "BOOLEAN", "Logical Flag", "1", "Active status indicator for dispatch assignment.", "1"),
        ]
    },
    {
        "num": 19,
        "name": "Aquaculture Species Data Table (STOCK_SPECIES)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique primary key for species taxonomy entry.", "1"),
            ("name", "VARCHAR", "Text", "100", "Common commercial aquaculture species name.", "Australian Redclaw Crayfish"),
            ("category", "VARCHAR", "Text", "20", "Broad farming classification (crayfish, superworm).", "crayfish"),
            ("scientific_name", "VARCHAR", "Text", "150", "Formal binomial biological nomenclature.", "Cherax quadricarinatus"),
            ("market_weight_g", "DECIMAL(8,2)", "Decimal", "8,2", "Target individual harvest weight in grams.", "65.00"),
            ("description", "TEXT", "Long text", "—", "Biological traits, feeding guide, and growth notes.", "Freshwater crustacean prized for high meat yield."),
        ]
    },
    {
        "num": 20,
        "name": "Ponds and Racks Data Table (OPERATIONS_POND)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique identifier for physical pond or rack unit.", "1"),
            ("name", "VARCHAR", "Text", "100", "Physical containment unit label or code.", "Concrete Pond A-1"),
            ("location", "VARCHAR", "Text", "255", "Physical zone location in Barangay Lantad facility.", "Zone 1 - Main Hatchery"),
            ("size", "DECIMAL(10,2)", "Decimal", "10,2", "Surface footprint area in square meters (sqm).", "24.00"),
            ("depth", "DECIMAL(5,2)", "Decimal", "5,2", "Water water column depth in meters.", "1.20"),
            ("capacity", "INT", "Number", "11", "Maximum biological stocking capacity count.", "500"),
            ("breeding_type", "VARCHAR", "Text", "50", "Operational role (NURSERY, GROWOUT, BREEDING).", "GROWOUT"),
            ("shelf_position", "VARCHAR", "Text", "50", "Rack tier positioning for superworm multi-tier trays.", "Tray Rack Tier 2"),
            ("status", "VARCHAR", "Text", "20", "Pond condition (ACTIVE, MAINTENANCE, FALLOW).", "ACTIVE"),
        ]
    },
    {
        "num": 21,
        "name": "Stock Batches Data Table (STOCK_STOCKBATCH)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique identifier for cultured stock batch.", "201"),
            ("pond_id (FK)", "INT", "Number", "11", "Foreign key referencing OPERATIONS_POND(id).", "1"),
            ("species_id (FK)", "INT", "Number", "11", "Foreign key referencing STOCK_SPECIES(id).", "1"),
            ("batch_code", "VARCHAR", "Alphanumeric", "50", "Unique tracking batch code format.", "BATCH-2025-CR-01"),
            ("stocking_date", "DATE", "Date", "—", "Calendar date when biological cohort was introduced.", "2025-02-01"),
            ("initial_quantity", "INT", "Number", "11", "Total individual juveniles or larvae stocked.", "400"),
            ("current_quantity", "INT", "Number", "11", "Current live stock headcount.", "385"),
            ("initial_avg_weight", "DECIMAL(8,2)", "Decimal", "8,2", "Initial individual weight in grams at stocking.", "5.50"),
            ("current_avg_weight", "DECIMAL(8,2)", "Decimal", "8,2", "Current sample weight in grams.", "52.30"),
            ("stage", "VARCHAR", "Text", "20", "Current life cycle stage (juvenile, sub_adult, adult).", "adult"),
            ("is_active", "BOOLEAN", "Logical Flag", "1", "Active status indicator for ongoing cultivation.", "1"),
        ]
    },
    {
        "num": 22,
        "name": "Feed Types Data Table (FEED_FEEDTYPE)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique primary key for feed formulation record.", "1"),
            ("name", "VARCHAR", "Text", "100", "Brand or commercial feed formulation name.", "Sinking Pellets (35% Protein)"),
            ("category", "VARCHAR", "Text", "20", "Feed class (crayfish_pellets, wheat_bran, carrots).", "crayfish_pellets"),
            ("accent_color", "VARCHAR", "Text", "20", "UI color token for dashboard analytics visualization.", "blue"),
            ("description", "TEXT", "Long text", "—", "Nutritional content, protein ratio, and instructions.", "High-protein sinking starter pellets for Redclaw."),
            ("is_active", "BOOLEAN", "Logical Flag", "1", "Status indicator for inventory reordering availability.", "1"),
        ]
    },
    {
        "num": 23,
        "name": "Feeding Logs Data Table (FEED_FEEDINGLOG)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique auto-incremented log transaction identifier.", "501"),
            ("stock_batch_id (FK)", "INT", "Number", "11", "Foreign key referencing STOCK_STOCKBATCH(id).", "201"),
            ("feed_type_id (FK)", "INT", "Number", "11", "Foreign key referencing FEED_FEEDTYPE(id).", "1"),
            ("quantity_kg", "DECIMAL(8,2)", "Decimal", "8,2", "Total mass of feed administered in kilograms.", "1.25"),
            ("feeding_time", "DATETIME", "Date & time", "—", "Timestamp when feeding regimen was carried out.", "2025-02-15 17:30:00"),
            ("fed_by_id (FK)", "INT", "Number", "11", "Foreign key referencing AUTH_USER(id) farm staff.", "1"),
        ]
    },
    {
        "num": 24,
        "name": "Harvest Schedules Data Table (HARVEST_HARVESTSCHEDULE)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique identifier for planned harvest forecast.", "1"),
            ("stock_batch_id (FK)", "INT", "Number", "11", "Foreign key referencing STOCK_STOCKBATCH(id).", "201"),
            ("scheduled_date", "DATE", "Date", "—", "Target scheduled harvest execution date.", "2025-06-15"),
            ("estimated_quantity", "INT", "Number", "11", "Projected headcount yield upon full maturity.", "380"),
            ("estimated_weight_kg", "DECIMAL(12,2)", "Decimal", "12,2", "Predicted gross harvest weight in kg.", "24.70"),
            ("target_weight_g", "DECIMAL(8,2)", "Decimal", "8,2", "Target minimum weight per specimen in grams.", "65.00"),
            ("status", "VARCHAR", "Text", "20", "Schedule status (PENDING, COMPLETED, CANCELLED).", "PENDING"),
        ]
    },
    {
        "num": 25,
        "name": "Harvest Records Data Table (HARVEST_HARVESTRECORD)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique identifier for actual harvest event record.", "301"),
            ("harvest_sched_id (FK)", "INT", "Number", "11", "Foreign key referencing HARVEST_HARVESTSCHEDULE(id).", "1"),
            ("stock_batch_id (FK)", "INT", "Number", "11", "Foreign key referencing STOCK_STOCKBATCH(id).", "201"),
            ("harvest_date", "DATE", "Date", "—", "Actual date when pond was harvested and sorted.", "2025-06-15"),
            ("quantity_harvested", "INT", "Number", "11", "Total live specimens harvested and retrieved.", "378"),
            ("total_weight_kg", "DECIMAL(12,2)", "Decimal", "12,2", "Total net weighed harvest yield in kilograms.", "25.10"),
            ("grade_a_quantity", "INT", "Number", "11", "Count of premium Grade A commercial specimens.", "280"),
            ("grade_b_quantity", "INT", "Number", "11", "Count of standard Grade B commercial specimens.", "75"),
            ("grade_c_quantity", "INT", "Number", "11", "Count of cull/sub-standard Grade C specimens.", "23"),
        ]
    },
    {
        "num": 26,
        "name": "Products and Pricing Data Table (SALES_PRODUCT)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique primary key for product catalog inventory.", "1"),
            ("species_id (FK)", "INT", "Number", "11", "Foreign key referencing STOCK_SPECIES(id).", "1"),
            ("pond_id (FK)", "INT", "Number", "11", "Foreign key referencing OPERATIONS_POND(id) source.", "1"),
            ("name", "VARCHAR", "Text", "200", "Commercial product catalog listing name.", "Live Australian Redclaw Crayfish (Grade A)"),
            ("category", "VARCHAR", "Text", "20", "Product species grouping (crayfish, superworm).", "crayfish"),
            ("unit_type", "VARCHAR", "Text", "20", "Pricing unit type (kg, pcs, tub, pair, pack).", "kg"),
            ("quantity_kg", "DECIMAL(12,2)", "Decimal", "12,2", "Currently available sellable stock weight in kg.", "45.50"),
            ("unit_price", "DECIMAL(12,2)", "Currency", "12,2", "Base sales retail selling price per unit in PHP.", "650.00"),
            ("pieces_per_kg", "DECIMAL(10,2)", "Decimal", "10,2", "Average specimen count per kilogram ratio.", "15.00"),
            ("placement_date", "DATE", "Date", "—", "Date inventory was made available on storefront.", "2025-06-16"),
        ]
    },
    {
        "num": 27,
        "name": "Sales Orders Data Table (SALES_SALESORDER)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique auto-incremented sales order identifier.", "1001"),
            ("order_number", "VARCHAR", "Alphanumeric", "50", "Unique generated transaction code.", "ORD-2025-06-001"),
            ("customer_id (FK)", "INT", "Number", "11", "Foreign key referencing SALES_CUSTOMER(id).", "101"),
            ("product_id (FK)", "INT", "Number", "11", "Foreign key referencing SALES_PRODUCT(id).", "1"),
            ("harvest_record_id (FK)", "INT", "Number", "11", "Optional foreign key linking to harvest origin.", "301"),
            ("order_date", "DATE", "Date", "—", "Date when buyer placed purchase transaction.", "2025-06-18"),
            ("delivery_date", "DATE", "Date", "—", "Requested or scheduled fulfillment date.", "2025-06-19"),
            ("quantity_kg", "DECIMAL(10,2)", "Decimal", "10,2", "Quantity of product purchased by the customer.", "10.00"),
            ("price_per_kg", "DECIMAL(10,2)", "Currency", "10,2", "Locked unit transaction price per unit/kg in PHP.", "650.00"),
            ("total_amount", "DECIMAL(12,2)", "Currency", "12,2", "Gross payable monetary order amount in PHP.", "6500.00"),
            ("status", "VARCHAR", "Text", "20", "Order state (PENDING, APPROVED, SHIPPED, DELIVERED).", "DELIVERED"),
            ("payment_status", "VARCHAR", "Text", "20", "Settlement state (UNPAID, PAID, GCASH_VERIFIED).", "PAID"),
        ]
    },
    {
        "num": 28,
        "name": "Order Deliveries Data Table (SALES_DELIVERY)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique delivery dispatch route identifier.", "1"),
            ("order_id (FK)", "INT", "Number", "11", "Foreign key referencing SALES_SALESORDER(id).", "1001"),
            ("rider_id (FK)", "INT", "Number", "11", "Foreign key referencing SALES_RIDER(id) courier.", "1"),
            ("scheduled_date", "DATE", "Date", "—", "Planned date for rider dispatch and delivery.", "2025-06-19"),
            ("delivered_date", "DATE", "Date", "—", "Actual completion timestamp of delivery.", "2025-06-19"),
            ("delivery_location", "VARCHAR", "Text", "255", "Full destination address with landmarks.", "Lacson St., Bacolod City"),
            ("quantity_kg", "DECIMAL(12,2)", "Decimal", "12,2", "Payload order delivery weight in kilograms.", "10.00"),
            ("status", "VARCHAR", "Text", "20", "Delivery state (ASSIGNED, IN_TRANSIT, DELIVERED).", "DELIVERED"),
        ]
    },
    {
        "num": 29,
        "name": "Inventory Transactions Data Table (SALES_INVENTORYTRANSACTION)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique inventory audit movement identifier.", "1"),
            ("product_id (FK)", "INT", "Number", "11", "Foreign key referencing SALES_PRODUCT(id).", "1"),
            ("related_order_id (FK)", "INT", "Number", "11", "Optional foreign key linking to sales deduction.", "1001"),
            ("related_harvest_id (FK)", "INT", "Number", "11", "Optional foreign key linking to harvest addition.", "301"),
            ("quantity_kg", "DECIMAL(12,2)", "Decimal", "12,2", "Stock change delta magnitude in kilograms.", "-10.00"),
            ("transaction_type", "VARCHAR", "Text", "20", "Movement type (HARVEST_IN, SALE_OUT, SPOILAGE).", "SALE_OUT"),
            ("created_by_id (FK)", "INT", "Number", "11", "Foreign key referencing AUTH_USER(id) initiator.", "1"),
            ("created_at", "DATETIME", "Date & time", "—", "Exact timestamp when stock ledger was recorded.", "2025-06-18 14:15:00"),
        ]
    },
    {
        "num": 30,
        "name": "System Audit and Input Logs Data Table (SALES_INPUTLOG)",
        "fields": [
            ("id (PK)", "INT (AI)", "Number", "11", "Unique auto-incremented security audit log entry.", "801"),
            ("user_id (FK)", "INT", "Number", "11", "Foreign key referencing AUTH_USER(id) actor.", "1"),
            ("user_name", "VARCHAR", "Text", "150", "Cached display username of the transaction actor.", "owner_silay"),
            ("user_role", "VARCHAR", "Text", "50", "Role tier at the moment of action execution.", "Farm Owner"),
            ("action", "VARCHAR", "Text", "20", "Executed database operation (CREATE, UPDATE, DELETE).", "CREATE"),
            ("module", "VARCHAR", "Text", "100", "System module accessed during event.", "Stock Management"),
            ("target_entity", "VARCHAR", "Text", "255", "Specific record identity modified or created.", "StockBatch: BATCH-2025-CR-01"),
            ("created_at", "DATETIME", "Date & time", "—", "Immutable system timestamp of event execution.", "2025-02-01 09:12:44"),
        ]
    }
]

def build_data_dictionary_docx():
    doc = Document()
    
    # Page Setup (8.5 x 11, Left 1.5, Top/Right/Bottom 1.0)
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1.0)
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)

    # Heading 3.7 Data Dictionary
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(12)
    p_h.paragraph_format.space_after = Pt(6)
    p_h.paragraph_format.line_spacing = 1.5
    r_h = p_h.add_run("3.7 Data Dictionary")
    r_h.font.name = "Arial"
    r_h.font.size = Pt(12)
    r_h.font.bold = True
    r_h.font.color.rgb = RGBColor(0, 0, 0)

    # Intro Paragraph
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.first_line_indent = Inches(0.5)
    r_intro = p_intro.add_run(
        "The data dictionary describes the structure, data types, formats, field sizes, descriptions, "
        "and example values for each table in the Lantad Superworm & Crayfish Management System database. "
        "The tables documented below represent the core data entities of the system."
    )
    r_intro.font.name = "Arial"
    r_intro.font.size = Pt(12)
    r_intro.font.color.rgb = RGBColor(0, 0, 0)

    # Build all 15 tables
    for item in DATA_DICTIONARY_TABLES:
        add_data_table(doc, item["num"], item["name"], item["fields"])

    out_path = os.path.abspath("Documents/Data_Dictionary.docx")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print(f"Data Dictionary DOCX saved successfully: {out_path}")

if __name__ == "__main__":
    build_data_dictionary_docx()

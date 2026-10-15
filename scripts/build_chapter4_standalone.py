import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
import sys
sys.path.insert(0, os.path.abspath('scripts'))
import build_chapter4_section4_1 as bc4
import build_chapter4_testing_tables as bc4t

def build_standalone_chapter4():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    # Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Arial'
    style_normal.font.size = Pt(12)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)

    # Chapter Header
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_head.paragraph_format.space_before = Pt(0)
    p_head.paragraph_format.space_after = Pt(12)
    p_head.paragraph_format.line_spacing = 1.5
    r = p_head.add_run("CHAPTER IV\nRESULTS AND DISCUSSION")
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True

    # Intro paragraphs
    p1 = doc.add_paragraph()
    p1.paragraph_format.first_line_indent = Inches(0.5)
    p1.paragraph_format.line_spacing = 1.5
    p1.paragraph_format.space_after = Pt(8)
    p1.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r1 = p1.add_run(
        "This chapter presents the results and discussions based on the development and evaluation of the "
        "Aquaculture Value Chain Management System with Predictive Analytics for Silay Superworm & Crayfish, in line "
        "with the objectives established in Chapter I. It includes comprehensive interface walkthroughs of the system's "
        "operational modules, the rigorous results of functionality testing (encompassing 68 Black-box functional test cases "
        "across twelve operational modules and 16 White-box structural test cases), and the software quality evaluation conducted in accordance with the "
        "ISO/IEC 25010:2011 Software Quality Model. Data were gathered from key stakeholders, including the Farm Owner, "
        "commercial buyers, retail customers, delivery riders, and independent IT experts, utilizing a five-point Likert scale "
        "ranging from 1 (Strongly Disagree) to 5 (Strongly Agree)."
    )
    r1.font.name = 'Arial'
    r1.font.size = Pt(12)

    p2 = doc.add_paragraph()
    p2.paragraph_format.first_line_indent = Inches(0.5)
    p2.paragraph_format.line_spacing = 1.5
    p2.paragraph_format.space_after = Pt(12)
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r2 = p2.add_run(
        "Furthermore, this chapter provides an empirical analysis of the system's execution under real-world aquaculture farm "
        "workflows in Barangay Lantad, Silay City, Negros Occidental. The findings demonstrate the system's technical effectiveness, "
        "operational reliability, data security, calculation precision in dual-unit pricing and inventory stock reconciliation, and capacity "
        "to eliminate manual recordkeeping inefficiencies while providing predictive decision-support capabilities through machine learning "
        "forecasting and customer segmentation."
    )
    r2.font.name = 'Arial'
    r2.font.size = Pt(12)

    # 4.1 Results of the Study
    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(14)
    p_h2.paragraph_format.space_after = Pt(6)
    p_h2.paragraph_format.line_spacing = 1.5
    r_h2 = p_h2.add_run("4.1 Results of the Study")
    r_h2.bold = True
    r_h2.font.name = 'Arial'
    r_h2.font.size = Pt(12)

    p3 = doc.add_paragraph()
    p3.paragraph_format.first_line_indent = Inches(0.5)
    p3.paragraph_format.line_spacing = 1.5
    p3.paragraph_format.space_after = Pt(8)
    p3.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r3 = p3.add_run(
        "The first objective of this study aimed to design and develop an Aquaculture Value Chain Management System with Predictive "
        "Analytics equipped with technical features including: (1) multi-role user authentication and session security, (2) an executive "
        "dashboard with real-time farm KPIs and sanitization alerts, (3) pond, cabin rack, and tank cultivation management, (4) feed formulation and daily "
        "supply logging, (5) Azula cultivation and 15-day pond sanitization monitoring, (6) harvest lifecycle scheduling and interactive production calendar, (7) "
        "commercial product catalog with dual-unit pricing (Per Piece and Per KG), (8) customer storefront with online cart checkout, (9) "
        "digital payment settlement supporting PayMongo GCash and COD, (10) sales order fulfillment and lifecycle state machine, (11) courier "
        "fleet dispatch and mobile rider portal with Proof of Delivery (POD), (12) real-time GPS tracking with Leaflet GIS maps, (13) predictive "
        "demand forecasting using Ordinary Least Squares Linear Regression, (14) customer segmentation using K-Means clustering, and (15) executive "
        "PDF/Excel report generation with a three-tiered audit trail."
    )
    r3.font.name = 'Arial'
    r3.font.size = Pt(12)

    p4 = doc.add_paragraph()
    p4.paragraph_format.first_line_indent = Inches(0.5)
    p4.paragraph_format.line_spacing = 1.5
    p4.paragraph_format.space_after = Pt(12)
    p4.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r4 = p4.add_run(
        "The development of the system was completed using an enterprise Model-View-Template (MVT) architecture combining Python 3.11 with the "
        "Django 5.1 framework on the backend, supported by an ACID-compliant PostgreSQL / SQLite relational database, responsive HTML5/CSS3/JavaScript "
        "frontend with Bootstrap 5 and Tailwind CSS styling, Chart.js/ApexCharts interactive visualization libraries, and Leaflet.js GIS spatial "
        "mapping. The operational interfaces of the developed system are presented in Figures 14 through 29."
    )
    r4.font.name = 'Arial'
    r4.font.size = Pt(12)

    for fig_num, fname, title, desc in bc4.CHAPTER_4_FIGURES:
        bc4.add_figure_item(doc, fig_num, title, desc, img_filename=fname)

    # 4.2 and 4.3 Testing and Evaluation
    bc4t.add_section_4_2_and_4_3_to_doc(doc)

    out_path = 'Documents/Final_Paper/Chapter_4.docx'
    doc.save(out_path)
    print(f"Successfully built {out_path}")

    downloads_path = r'C:\Users\kymer\Downloads\Chapter_4 redo.docx'
    try:
        doc.save(downloads_path)
        print(f"Successfully updated {downloads_path}")
    except Exception as e:
        print(f"Warning: Could not save to {downloads_path}: {e}")

if __name__ == '__main__':
    build_standalone_chapter4()

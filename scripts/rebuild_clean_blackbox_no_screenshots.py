import os
import io
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in tcPr.findall(qn('w:shd')):
        tcPr.remove(child)
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=60, bottom=60, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in tcPr.findall(qn('w:tcMar')):
        tcPr.remove(child)
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_table_borders(table, color="000000", sz="4"):
    tblPr = table._tbl.tblPr
    for child in tblPr.findall(qn('w:tblBorders')):
        tblPr.remove(child)
    borders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>'
        f'<w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:right w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'<w:insideV w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/>'
        f'</w:tblBorders>'
    )
    tblPr.append(borders)

def make_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="1"/>'))

def make_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")} w:val="1"/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")} w:val="1"/>'))

def format_paragraph(p, text, font_name="Arial", font_size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6, color=(0, 0, 0)):
    p.text = ""
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.name = font_name
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    return run

def add_table_title(doc, table_num, title_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.0
    
    r1 = p.add_run(f"Table {table_num}\n")
    r1.font.name = "Arial"
    r1.font.size = Pt(11)
    r1.font.bold = True
    r1.font.color.rgb = RGBColor(0, 0, 0)
    
    r2 = p.add_run(title_text)
    r2.font.name = "Arial"
    r2.font.size = Pt(11)
    r2.font.italic = True
    r2.font.color.rgb = RGBColor(0, 0, 0)

def rebuild_clean_blackbox():
    src_path = os.path.abspath("Documents/Chapter 4/blackbox final.docx")
    target_path = os.path.abspath("Documents/Chapter 4/blackbox final.docx")
    doc_src = Document(src_path)

    # 1. Extract test cases from the 11 Black-box tables (excluding merged drawing rows)
    raw_modules = []
    for t_idx in range(11):
        tbl = doc_src.tables[t_idx]
        tcs = []
        for r_idx in range(1, len(tbl.rows)):
            row = tbl.rows[r_idx]
            txt = row.cells[0].text.strip()
            if txt.startswith("Evidence:"):
                continue  # Skip old evidence image row
            tc_id = row.cells[0].text.strip()
            desc = row.cells[1].text.strip()
            action = row.cells[2].text.strip()
            inp = row.cells[3].text.strip()
            actual = row.cells[4].text.strip()
            expected = row.cells[5].text.strip()
            pf = row.cells[6].text.strip()
            tcs.append({
                'orig_id': tc_id,
                'desc': desc,
                'action': action,
                'input': inp,
                'actual': actual,
                'expected': expected,
                'pass_fail': pf
            })
        raw_modules.append(tcs)

    module_meta = [
        {"num": 52, "name": "Registration, Authentication, and Session Security", "prefix": "AUTHENTICATE"},
        {"num": 53, "name": "Executive Dashboard and Farm Operations Monitoring", "prefix": "DASHBOARD"},
        {"num": 54, "name": "Commercial Product Inventory Management", "prefix": "PRODUCT"},
        {"num": 55, "name": "Aquaculture Farm Operations (Ponds, Superworm Cabin, and Azula Cultivation)", "prefix": "AQUACULTURE"},
        {"num": 56, "name": "Harvest Lifecycle Management and Production Calendar", "prefix": "HARVEST"},
        {"num": 57, "name": "Customer Storefront Cart and Fulfillment", "prefix": "STOREFRONT"},
        {"num": 58, "name": "Sales Orders and Purchases Management", "prefix": "ORDER"},
        {"num": 59, "name": "Digital Payment Settlement and PayMongo Gateway", "prefix": "PAYMENT"},
        {"num": 60, "name": "Delivery Fleet and Courier Logistics", "prefix": "FLEET"},
        {"num": 61, "name": "Real-Time GPS Order Tracking and Telemetry", "prefix": "TRACKING"},
        {"num": 62, "name": "Predictive Demand Analytics and Reporting", "prefix": "ANALYTICS"}
    ]

    new_doc = Document()
    sec = new_doc.sections[0]
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.0)
    sec.right_margin = Inches(1.0)
    sec.page_width = Inches(8.5)
    sec.page_height = Inches(11.0)

    # Document Header
    p_head = new_doc.add_paragraph()
    format_paragraph(p_head, "CHAPTER IV\nRESULTS AND DISCUSSION", font_size=14, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)

    p_intro = new_doc.add_paragraph()
    format_paragraph(
        p_intro,
        "The second objective of this study aimed to evaluate the technical functionality, security controls, biological lifecycle tracking, real-time dispatch, and data integrity of the Lantad Superworm & Crawfish Farm Management and Online Marketplace System through comprehensive Black-box Alpha Testing. Test cases were structured to validate all operational pathways across owner, administrator, delivery rider, and commercial customer roles under both deterministic and boundary conditions.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=8
    )

    p_intro2 = new_doc.add_paragraph()
    format_paragraph(
        p_intro2,
        "A total of 66 curated, high-impact Black-box test cases were formulated, executed, and verified across eleven (11) core functional modules. The testing execution adhered strictly to structured verification procedures with unique module identifiers, presented in the clean Table-then-Evidence format.",
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
        space_after=12
    )

    col_widths = [Inches(1.20), Inches(0.95), Inches(1.05), Inches(0.90), Inches(1.00), Inches(1.00), Inches(0.40)]
    headers = ['Test Case ID', 'Description', 'Action', 'Input Data', 'Actual Outcome', 'Expected Outcome', 'Pass/Fail']

    # Render each module: Table FIRST, then Evidence format underneath
    for m_idx, meta in enumerate(module_meta):
        tcs = raw_modules[m_idx]
        title_text = f"Black-Box Testing of {meta['name']} (Cycle 1 & Cycle 2)"
        add_table_title(new_doc, meta['num'], title_text)
        
        p_desc = new_doc.add_paragraph()
        first_id = tcs[0]['orig_id'] if tcs else 'TC-001'
        last_id = tcs[-1]['orig_id'] if tcs else 'TC-006'
        format_paragraph(
            p_desc,
            f"Table {meta['num']} summarizes the {meta['name'].lower()} test cases ({first_id} through {last_id}) during Alpha Testing. The evaluation verified all functional requirements, role permissions, state transitions, and edge cases across {len(tcs)} test cases.",
            align=WD_ALIGN_PARAGRAPH.JUSTIFY,
            space_after=6
        )

        # 1. PURE CLEAN TEST CASE TABLE (No image drawings inside the table)
        tbl = new_doc.add_table(rows=len(tcs) + 1, cols=7)
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(tbl, color="000000", sz="4")

        # Header Row
        hdr_row = tbl.rows[0]
        make_row_header(hdr_row)
        for c_idx, cell in enumerate(hdr_row.cells):
            cell.width = col_widths[c_idx]
            set_cell_background(cell, "F2F4F8")
            set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            p = cell.paragraphs[0]
            format_paragraph(
                p,
                headers[c_idx],
                font_size=10,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER if c_idx in [0, 6] else WD_ALIGN_PARAGRAPH.LEFT,
                space_after=0
            )

        # Data Rows
        for r_i, tc in enumerate(tcs):
            row = tbl.rows[r_i + 1]
            make_row_cant_split(row)
            row_vals = [tc['orig_id'], tc['desc'], tc['action'], tc['input'], tc['actual'], tc['expected'], tc['pass_fail']]
            for c_idx, cell in enumerate(row.cells):
                cell.width = col_widths[c_idx]
                set_cell_margins(cell, top=50, bottom=50, left=70, right=70)
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                p = cell.paragraphs[0]
                is_centered = c_idx in [0, 6]
                is_bold = c_idx in [0, 6]
                format_paragraph(
                    p,
                    row_vals[c_idx],
                    font_size=8.5 if c_idx == 0 else 9,
                    bold=is_bold,
                    align=WD_ALIGN_PARAGRAPH.CENTER if is_centered else WD_ALIGN_PARAGRAPH.LEFT,
                    space_after=0
                )

        # 2. EVIDENCE FORMAT DIRECTLY UNDERNEATH THE TABLE (Clean evidence blocks ready for screenshots)
        p_ev_header = new_doc.add_paragraph()
        format_paragraph(p_ev_header, f"Test Verification Evidence for Table {meta['num']} ({meta['name']})", font_size=11, bold=True, space_before=10, space_after=4)

        for tc in tcs:
            # Evidence Caption
            p_cap = new_doc.add_paragraph()
            format_paragraph(
                p_cap,
                f"Evidence: {tc['orig_id']} - {tc['desc']}",
                font_size=10,
                bold=True,
                align=WD_ALIGN_PARAGRAPH.CENTER,
                space_before=6,
                space_after=2
            )
            
            # Evidence Placeholder Box (Clean 1x1 table placeholder with no embedded drawing image)
            t_box = new_doc.add_table(rows=1, cols=1)
            t_box.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(t_box, color="CCCCCC", sz="4")
            c_box = t_box.rows[0].cells[0]
            c_box.width = Inches(6.5)
            set_cell_background(c_box, "FAFAFA")
            set_cell_margins(c_box, top=140, bottom=140, left=100, right=100)
            p_box = c_box.paragraphs[0]
            format_paragraph(
                p_box,
                f"[ Screenshot Evidence Placeholder for {tc['orig_id']} ]",
                font_size=9.5,
                italic=True,
                color=(120, 120, 120),
                align=WD_ALIGN_PARAGRAPH.CENTER,
                space_after=0
            )

        new_doc.add_paragraph()  # Spacing between suites

    # Save output cleanly
    new_doc.save(target_path)
    print(f"Successfully rebuilt clean blackbox document without heavy screenshots: {target_path}")

if __name__ == '__main__':
    rebuild_clean_blackbox()

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_apa_table_borders(table):
    tblPr = table._element.xpath('w:tblPr')
    if tblPr:
        borders = parse_xml(
            f'<w:tblBorders {nsdecls("w")}>\n'
            f'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
            f'  <w:left w:val="none"/>\n'
            f'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>\n'
            f'  <w:right w:val="none"/>\n'
            f'  <w:insideH w:val="none"/>\n'
            f'  <w:insideV w:val="none"/>\n'
            f'</w:tblBorders>'
        )
        tblPr[0].append(borders)

def set_row_bottom_border(row):
    for cell in row.cells:
        tcPr = cell._element.get_or_add_tcPr()
        tcBorders = parse_xml(
            f'<w:tcBorders {nsdecls("w")}>\n'
            f'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
            f'</w:tcBorders>'
        )
        tcPr.append(tcBorders)

def add_custom_table(doc, table_num, table_name, headers, col_widths, rows_data):
    p_num = doc.add_paragraph()
    p_num.paragraph_format.space_before = Pt(14)
    p_num.paragraph_format.space_after = Pt(2)
    p_num.paragraph_format.line_spacing = 1.0
    r_num = p_num.add_run(f"Table {table_num}")
    r_num.font.name = "Arial"
    r_num.font.size = Pt(11)
    r_num.font.bold = True
    r_num.font.color.rgb = RGBColor(0, 0, 0)

    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    p_title.paragraph_format.line_spacing = 1.0
    r_title = p_title.add_run(table_name)
    r_title.font.name = "Arial"
    r_title.font.size = Pt(11)
    r_title.font.bold = False
    r_title.font.color.rgb = RGBColor(0, 0, 0)

    table = doc.add_table(rows=len(rows_data) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_apa_table_borders(table)

    # Header
    hdr_row = table.rows[0]
    set_row_bottom_border(hdr_row)
    for idx, heading in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.width = col_widths[idx]
        set_cell_margins(cell, top=80, bottom=80, left=100, right=100)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(heading)
        r.font.name = "Arial"
        r.font.size = Pt(10)
        r.font.bold = True
        r.font.color.rgb = RGBColor(0, 0, 0)

    # Data Rows
    for row_idx, data in enumerate(rows_data):
        row = table.rows[row_idx + 1]
        for col_idx, text_val in enumerate(data):
            cell = row.cells[col_idx]
            cell.width = col_widths[col_idx]
            set_cell_margins(cell, top=60, bottom=60, left=100, right=100)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
            r = p.add_run(str(text_val))
            r.font.name = "Arial"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0, 0, 0)

    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(8)
    p_sp.paragraph_format.line_spacing = 1.0

LIKERT_SCALE_ROWS = [
    ("5", "4.21 – 5.00", "Strongly Agree (SA)", "Excellent / Highly Acceptable"),
    ("4", "3.41 – 4.20", "Agree (A)", "Very Good / Acceptable"),
    ("3", "2.61 – 3.40", "Moderately Agree (MA)", "Good / Moderately Acceptable"),
    ("2", "1.81 – 2.60", "Disagree (D)", "Fair / Needs Improvement"),
    ("1", "1.00 – 1.80", "Strongly Disagree (SD)", "Poor / Unacceptable")
]

def add_section_3_10(doc):
    # Heading: 3.10 System Evaluation Tools
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(14)
    p_h.paragraph_format.space_after = Pt(6)
    p_h.paragraph_format.line_spacing = 1.5
    r_h = p_h.add_run("3.10 System Evaluation Tools")
    r_h.font.name = "Arial"
    r_h.font.size = Pt(12)
    r_h.font.bold = True
    r_h.font.color.rgb = RGBColor(0, 0, 0)

    # Subheading: White-Box Testing
    p_sub1 = doc.add_paragraph()
    p_sub1.paragraph_format.space_before = Pt(6)
    p_sub1.paragraph_format.space_after = Pt(4)
    p_sub1.paragraph_format.line_spacing = 1.5
    r_sub1 = p_sub1.add_run("White-Box Testing")
    r_sub1.font.name = "Arial"
    r_sub1.font.size = Pt(12)
    r_sub1.font.bold = True
    r_sub1.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph 1
    p_wb1 = doc.add_paragraph()
    p_wb1.paragraph_format.space_before = Pt(0)
    p_wb1.paragraph_format.space_after = Pt(8)
    p_wb1.paragraph_format.line_spacing = 1.5
    p_wb1.paragraph_format.first_line_indent = Inches(0.5)
    r_wb1 = p_wb1.add_run(
        "White-Box Testing was conducted to evaluate the internal logic, backend processes, database operations, "
        "and security-related functions of the Aquaculture Value Chain Management System with Predictive Analytics for "
        "Silay Superworm & Crayfish. This testing approach used the Django Test Framework on the Python backend to examine "
        "the system's code structure and internal workflows and confirm that each module performed according to its intended function."
    )
    r_wb1.font.name = "Arial"
    r_wb1.font.size = Pt(12)
    r_wb1.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph 2
    p_wb2 = doc.add_paragraph()
    p_wb2.paragraph_format.space_before = Pt(0)
    p_wb2.paragraph_format.space_after = Pt(10)
    p_wb2.paragraph_format.line_spacing = 1.5
    p_wb2.paragraph_format.first_line_indent = Inches(0.5)
    r_wb2 = p_wb2.add_run(
        "The researchers used White-Box Testing to validate important backend processes such as user authentication, "
        "role-based access control, pond batch allocation, feeding and mortality record computation, harvest record processing, "
        "product inventory deduction, pricing calculation (per piece and per kilogram), sales order state transitions, "
        "payment verification, delivery dispatch assignment, Proof of Delivery validation, Ordinary Least Squares Linear Regression "
        "forecasting formulas, K-Means clustering routines, audit trail logging, and database backup operations. Through this testing, "
        "the researchers determined whether the system correctly processed data, applied validation rules, stored records properly, "
        "and maintained relational data integrity across all operational modules."
    )
    r_wb2.font.name = "Arial"
    r_wb2.font.size = Pt(12)
    r_wb2.font.color.rgb = RGBColor(0, 0, 0)

    # Subheading: Black-Box Testing
    p_sub2 = doc.add_paragraph()
    p_sub2.paragraph_format.space_before = Pt(8)
    p_sub2.paragraph_format.space_after = Pt(4)
    p_sub2.paragraph_format.line_spacing = 1.5
    r_sub2 = p_sub2.add_run("Black-Box Testing")
    r_sub2.font.name = "Arial"
    r_sub2.font.size = Pt(12)
    r_sub2.font.bold = True
    r_sub2.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph
    p_bb = doc.add_paragraph()
    p_bb.paragraph_format.space_before = Pt(0)
    p_bb.paragraph_format.space_after = Pt(10)
    p_bb.paragraph_format.line_spacing = 1.5
    p_bb.paragraph_format.first_line_indent = Inches(0.5)
    r_bb = p_bb.add_run(
        "Black-Box Testing was performed to validate the functional behavior and operational features of the system from the "
        "end-users' perspective without inspecting the internal source code. This testing methodology focused on testing user inputs, "
        "graphical interface responsiveness, navigation flows, and output generation across the Farm Owner, Commercial Customer, "
        "and Delivery Rider interfaces. Test cases were structured to verify critical features including user account registration "
        "and login, dynamic overview dashboard widgets, pond and stock batch management, feeding and mortality logging, commercial "
        "product catalog management, customer online shopping cart checkout, order payment settlement, rider dispatch routing, "
        "Proof of Delivery image upload, executive PDF and Excel report generation, and predictive analytics visualization."
    )
    r_bb.font.name = "Arial"
    r_bb.font.size = Pt(12)
    r_bb.font.color.rgb = RGBColor(0, 0, 0)

    # Subheading: System Software Quality Evaluation Instrument
    p_sub3 = doc.add_paragraph()
    p_sub3.paragraph_format.space_before = Pt(8)
    p_sub3.paragraph_format.space_after = Pt(4)
    p_sub3.paragraph_format.line_spacing = 1.5
    r_sub3 = p_sub3.add_run("System Software Quality Evaluation Instrument")
    r_sub3.font.name = "Arial"
    r_sub3.font.size = Pt(12)
    r_sub3.font.bold = True
    r_sub3.font.color.rgb = RGBColor(0, 0, 0)

    # Paragraph
    p_sq = doc.add_paragraph()
    p_sq.paragraph_format.space_before = Pt(0)
    p_sq.paragraph_format.space_after = Pt(8)
    p_sq.paragraph_format.line_spacing = 1.5
    p_sq.paragraph_format.first_line_indent = Inches(0.5)
    r_sq = p_sq.add_run(
        "The ISO/IEC 25010:2011 Software Quality Model will be used as the primary instrument to evaluate the quality of the "
        "developed system based on eight characteristics: Functional Suitability, Performance Efficiency, Compatibility, Usability, "
        "Reliability, Security, Maintainability, and Portability. A 5-point Likert scale will be utilized to gather quantitative feedback "
        "from respondents, who will assess how well the system meets each quality criterion. The collected results will be analyzed to "
        "determine the overall system quality and identify areas for improvement prior to its final deployment."
    )
    r_sq.font.name = "Arial"
    r_sq.font.size = Pt(12)
    r_sq.font.color.rgb = RGBColor(0, 0, 0)

    # Table 33
    add_custom_table(
        doc,
        table_num=33,
        table_name="5-Point Likert Scale, Mean Range, and Verbal Interpretation for ISO/IEC 25010:2011 System Software Quality Evaluation",
        headers=["Option / Rating", "Mean Range", "Verbal Description", "Qualitative Interpretation"],
        col_widths=[Inches(1.2), Inches(1.3), Inches(2.2), Inches(2.3)],
        rows_data=LIKERT_SCALE_ROWS
    )

def build_standalone():
    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    add_section_3_10(doc)
    out_path = os.path.abspath("Documents/System_Evaluation_Tools_3_10.docx")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print(f"Saved: {out_path}")

if __name__ == "__main__":
    build_standalone()

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import sys
import os

sys.path.insert(0, os.path.abspath('scripts'))
import use_cases_data

def add_use_case_table_to_doc(doc, uc):
    # 1. Table Title (e.g. Table 3 Authentication, Frictionless Phone Login & Role-Based Access Control)
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    p_title.paragraph_format.line_spacing = 1.15
    r_t = p_title.add_run(f"Table {uc['table_num']} {uc['title']}")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(12)
    r_t.bold = True

    # 2. 1-Column Table (2 rows: Row 0 is Use Case Title, Row 1 is Content)
    table = doc.add_table(rows=2, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.0)

    # Set APA-style single borders (top, bottom, inside horizontal border between header and body, no left/right/vertical)
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:left w:val="nil"/>\n'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:right w:val="nil"/>\n'
        f'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:insideV w:val="nil"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # Row 0: Use Case Header (e.g. Use Case 1: Authentication, Frictionless Phone Login & Role-Based Access Control)
    cell_0 = table.rows[0].cells[0]
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(4)
    p0.paragraph_format.space_after = Pt(4)
    p0.paragraph_format.line_spacing = 1.15
    r0 = p0.add_run(uc['uc_id'])
    r0.font.name = "Arial"
    r0.font.size = Pt(11)
    r0.bold = True

    # Row 1: Content
    cell_1 = table.rows[1].cells[0]
    
    # Goal paragraph
    p_goal = cell_1.paragraphs[0]
    p_goal.paragraph_format.space_before = Pt(4)
    p_goal.paragraph_format.space_after = Pt(3)
    p_goal.paragraph_format.line_spacing = 1.15
    rg1 = p_goal.add_run("Use Case Goal: ")
    rg1.bold = True
    rg1.font.name = "Arial"
    rg1.font.size = Pt(10)
    rg2 = p_goal.add_run(uc['goal'])
    rg2.font.name = "Arial"
    rg2.font.size = Pt(10)

    for sc_idx, sc in enumerate(uc['scenarios']):
        # Scenario Title
        p_sc = cell_1.add_paragraph()
        p_sc.paragraph_format.space_before = Pt(6 if sc_idx > 0 else 3)
        p_sc.paragraph_format.space_after = Pt(1)
        p_sc.paragraph_format.line_spacing = 1.15
        r_sc = p_sc.add_run(sc['name'])
        r_sc.bold = True
        r_sc.font.name = "Arial"
        r_sc.font.size = Pt(10)

        # Actor
        p_act = cell_1.add_paragraph()
        p_act.paragraph_format.space_before = Pt(0)
        p_act.paragraph_format.space_after = Pt(1)
        p_act.paragraph_format.line_spacing = 1.15
        r_lbl = p_act.add_run("Actor: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_val = p_act.add_run(sc['actor'])
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10)

        # Pre-condition
        p_pre = cell_1.add_paragraph()
        p_pre.paragraph_format.space_before = Pt(0)
        p_pre.paragraph_format.space_after = Pt(1)
        p_pre.paragraph_format.line_spacing = 1.15
        r_lbl = p_pre.add_run("Pre-condition: ")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)
        r_val = p_pre.add_run(sc.get('pre', ''))
        r_val.font.name = "Arial"
        r_val.font.size = Pt(10)

        # Main Flow
        p_mf = cell_1.add_paragraph()
        p_mf.paragraph_format.space_before = Pt(0)
        p_mf.paragraph_format.space_after = Pt(1)
        p_mf.paragraph_format.line_spacing = 1.15
        r_lbl = p_mf.add_run("Main Flow:")
        r_lbl.bold = True
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(10)

        # Steps
        for step_idx, step in enumerate(sc['steps'], 1):
            p_st = cell_1.add_paragraph()
            p_st.paragraph_format.space_before = Pt(0)
            p_st.paragraph_format.space_after = Pt(1)
            p_st.paragraph_format.line_spacing = 1.15
            
            # Format step text
            if step.startswith('Step '):
                parts = step.split(':', 1)
                r_st_lbl = p_st.add_run(parts[0] + ':')
                r_st_lbl.bold = True
                r_st_lbl.font.name = "Arial"
                r_st_lbl.font.size = Pt(10)
                r_st_val = p_st.add_run(parts[1])
                r_st_val.font.name = "Arial"
                r_st_val.font.size = Pt(10)
            else:
                r_st_lbl = p_st.add_run(f"Step {step_idx}: ")
                r_st_lbl.bold = True
                r_st_lbl.font.name = "Arial"
                r_st_lbl.font.size = Pt(10)
                r_st_val = p_st.add_run(step)
                r_st_val.font.name = "Arial"
                r_st_val.font.size = Pt(10)

        # User Action & System Response pairs
        for action, response in sc.get('pairs', []):
            p_ua = cell_1.add_paragraph()
            p_ua.paragraph_format.space_before = Pt(0)
            p_ua.paragraph_format.space_after = Pt(1)
            p_ua.paragraph_format.line_spacing = 1.15
            r_ua_lbl = p_ua.add_run("User Action: ")
            r_ua_lbl.bold = True
            r_ua_lbl.font.name = "Arial"
            r_ua_lbl.font.size = Pt(10)
            r_ua_val = p_ua.add_run(action)
            r_ua_val.font.name = "Arial"
            r_ua_val.font.size = Pt(10)

            p_sr = cell_1.add_paragraph()
            p_sr.paragraph_format.space_before = Pt(0)
            p_sr.paragraph_format.space_after = Pt(2)
            p_sr.paragraph_format.line_spacing = 1.15
            r_sr_lbl = p_sr.add_run("System Response: ")
            r_sr_lbl.bold = True
            r_sr_lbl.font.name = "Arial"
            r_sr_lbl.font.size = Pt(10)
            r_sr_val = p_sr.add_run(response)
            r_sr_val.font.name = "Arial"
            r_sr_val.font.size = Pt(10)

    # Spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
    p_sp.paragraph_format.line_spacing = 1.0

def build_standalone_use_case_doc():
    doc = docx.Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    p_head = doc.add_paragraph()
    p_head.paragraph_format.space_before = Pt(0)
    p_head.paragraph_format.space_after = Pt(6)
    p_head.paragraph_format.line_spacing = 1.5
    r_h = p_head.add_run("3.6 Use Case Scenario")
    r_h.font.name = "Arial"
    r_h.font.size = Pt(12)
    r_h.bold = True

    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(6)
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.first_line_indent = Inches(0.5)
    r_i = p_intro.add_run(
        "The Use Case Scenario describes the interaction between the primary entities (Admin / Farm Owner, "
        "Delivery Rider, and Commercial Customers / Buyers) and the Silay Superworm & Crayfish Value Chain "
        "Management System with Predictive Analytics. The Admin / Farm Owner and Delivery Rider authenticate "
        "through a unified system login interface (/login/) with strict role-based routing, while Commercial "
        "Customers register and log in via a dedicated, frictionless phone-based verification workflow (/customer/login/ "
        "and storefront modal). Each scenario specifies operational pre-conditions, step-by-step main flows of activities, "
        "actor actions, and corresponding system responses across farm operations, 15-day Azula pond sanitization schedules, "
        "cabin shelf grids, dual pricing (per-piece and per-kg), digital GCash online payment gateways, mobile courier dispatching "
        "with camera-verified Proof of Delivery (POD), role-segregated mobile-responsive notification feeds, 7-day machine learning "
        "sales forecasting, and comprehensive three-tiered audit and printable report documentation."
    )
    r_i.font.name = "Arial"
    r_i.font.size = Pt(12)

    for uc in use_cases_data.use_cases:
        add_use_case_table_to_doc(doc, uc)

    doc.save("Documents/Use_Case_Scenario.docx")
    print("Successfully built Documents/Use_Case_Scenario.docx with exact Table formatting!")

if __name__ == '__main__':
    build_standalone_use_case_doc()

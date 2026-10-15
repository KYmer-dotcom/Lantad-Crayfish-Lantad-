import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Insert scripts folder to sys.path
sys.path.insert(0, os.path.abspath('scripts'))
import build_chapter4_section4_1 as bc4_sec1
import build_chapter4_testing_tables as bc4_testing

def p_chapter_head(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(14)
    r.font.bold = True
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def add_body_p(doc, text, space_after=8, indent=0.5):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.5
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if indent > 0:
        p.paragraph_format.first_line_indent = Inches(indent)
    r = p.add_run(text)
    r.font.name = 'Arial'
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0, 0, 0)
    return p

def build_full_chapter4():
    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    # Chapter 4 Heading
    p_chapter_head(doc, "CHAPTER 4\nRESULTS AND DISCUSSION")

    # Introductory narrative
    add_body_p(
        doc,
        "This chapter presents the empirical results, structural implementations, software verification outcomes, "
        "and quality evaluation findings of the developed Aquaculture Value Chain Management System with Predictive "
        "Analytics for Silay Superworm & Crayfish Farm in Barangay Lantad, Silay City, Negros Occidental. The presentation "
        "of results is organized sequentially across three core sections corresponding to the specific objectives of the study: "
        "(1) the architectural results and user interface implementations of the developed system modules, (2) the technical "
        "verification outcomes from comprehensive Black-box and White-box testing, and (3) the software quality evaluation results "
        "based on the ISO/IEC 25010:2011 Software Product Quality Model."
    )

    # 4.1 Results of the Study
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(14)
    p_h1.paragraph_format.space_after = Pt(6)
    p_h1.paragraph_format.line_spacing = 1.5
    r_h1 = p_h1.add_run("4.1 Results of the Study")
    r_h1.bold = True
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(12)
    r_h1.font.color.rgb = RGBColor(0, 0, 0)

    add_body_p(
        doc,
        "The first objective of this study aimed to design and develop an Aquaculture Value Chain Management System with Predictive Analytics equipped "
        "with technical features including: (1) multi-role user authentication and session security, (2) an executive dashboard with real-time "
        "farm KPIs, (3) pond, cabin rack, and tank cultivation management, (4) feed formulation and daily feeding logging, (5) mortality tracking and survival "
        "rate computation, (6) harvest scheduling and commercial grade sorting, (7) commercial product catalog with dual-unit pricing (Per Piece and Per KG), "
        "(8) customer storefront with online cart checkout, (9) digital payment settlement supporting PayMongo GCash and COD, (10) sales order fulfillment "
        "and lifecycle state machine, (11) courier fleet dispatch and mobile rider portal with Proof of Delivery (POD), (12) real-time GPS tracking with Leaflet GIS maps, "
        "(13) predictive demand forecasting using Ordinary Least Squares Linear Regression, (14) customer segmentation using K-Means clustering, and (15) executive "
        "PDF/Excel report generation with a three-tiered audit trail."
    )

    add_body_p(
        doc,
        "The development of the system was completed using an enterprise Model-View-Template (MVT) architecture combining Python 3.11 with the Django 5.1 "
        "framework on the backend, supported by an ACID-compliant PostgreSQL / SQLite relational database, responsive HTML5/CSS3/JavaScript frontend with "
        "Bootstrap 5 and Tailwind CSS styling, Chart.js/ApexCharts interactive visualization libraries, and Leaflet.js GIS spatial mapping. The operational "
        "interfaces of the developed system are presented in Figures 14 through 29."
    )

    # Insert Figures 14 through 29
    for fig_num, fname, title, desc in bc4_sec1.CHAPTER_4_FIGURES:
        bc4_sec1.add_figure_item(doc, fig_num, title, desc, img_filename=fname)

    # 4.2 Black-box and White-box Testing Results (Tables 34 to 47)
    bc4_testing.add_section_4_2_to_doc(doc)

    out_ch4 = os.path.abspath("Documents/Chapter 4/Chapter_4_Results_and_Discussions.docx")
    os.makedirs(os.path.dirname(out_ch4), exist_ok=True)
    doc.save(out_ch4)
    print("Saved complete Chapter 4 to:", out_ch4)
    
    out_root = os.path.abspath("Documents/Chapter 4 - Results and Discussions.docx")
    doc.save(out_root)
    print("Saved complete Chapter 4 to root Documents:", out_root)

if __name__ == "__main__":
    build_full_chapter4()

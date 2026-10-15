import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_chapter_1():
    doc = docx.Document()

    # Configure Short Bond Paper (Letter: 8.5" x 11.0") with 1.5" left binding margin
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    # Base Normal Style: Arial 12pt, 1.5 line spacing, 0 space before, 6pt space after
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Arial'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)
    normal_style.paragraph_format.line_spacing = 1.5
    normal_style.paragraph_format.space_before = Pt(0)
    normal_style.paragraph_format.space_after = Pt(6)

    # Helper functions
    def p_title(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_center(text, bold=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.bold = bold
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_blank(count=1):
        for _ in range(count):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.5

    def p_chapter_head(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_heading_3(text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_body(text, indent=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if indent:
            p.paragraph_format.first_line_indent = Inches(0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        run = p.add_run(text)
        run.font.name = 'Arial'
        run.font.size = Pt(12)
        return p

    def p_term_item(term_bold, def_text):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.5
        run_bold = p.add_run(term_bold)
        run_bold.bold = True
        run_bold.font.name = 'Arial'
        run_bold.font.size = Pt(12)
        run_body = p.add_run(def_text)
        run_body.font.name = 'Arial'
        run_body.font.size = Pt(12)
        return p

    def p_objective_item(num_prefix, text, sub=False):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.5
        if sub:
            p.paragraph_format.left_indent = Inches(0.5)
        run_num = p.add_run(num_prefix)
        run_num.font.name = 'Arial'
        run_num.font.size = Pt(12)
        run_text = p.add_run(text)
        run_text.font.name = 'Arial'
        run_text.font.size = Pt(12)
        return p

    # ==================== TITLE PAGE ====================
    p_title("AQUACULTURE VALUE CHAIN MANAGEMENT SYSTEM WITH PREDICTIVE ANALYTICS")
    
    p_blank(3)
    
    p_center("A Capstone Project")
    p_center("Presented to the Faculty of the")
    p_center("College of Computer Studies")
    p_center("Carlos Hilado Memorial State University")
    
    p_blank(3)
    
    p_center("In Partial Fulfilment")
    p_center("of the Requirements for the Degree")
    p_center("BACHELOR OF SCIENCE IN INFORMATION SYSTEMS", bold=True)
    
    p_blank(3)
    
    p_center("By")
    p_blank(1)
    p_center("RESEARCH PROPONENTS", bold=True)
    
    p_blank(3)
    
    p_center("OCTOBER 2026", bold=True)
    
    # Page Break to Chapter 1
    doc.add_page_break()

    # ==================== CHAPTER I ====================
    p_chapter_head("CHAPTER I")
    p_chapter_head("Introduction")

    p_body("Superworm and crayfish production play an emerging role in sustainable food production, animal feed development, and livelihood generation, particularly in small-scale and rural farming communities. Superworms (Zophobas morio) are increasingly recognized as a promising insect-based protein source for animal feed and other nutritional applications, while crayfish, particularly freshwater species such as redclaw crayfish (Cherax quadricarinatus), are cultivated as high-value freshwater crustaceans in aquaculture systems (Rumbos & Athanassiou, 2021; Rigg et al., 2020). As demand for alternative protein sources and freshwater aquatic products continues to rise, the expansion of superworm and crayfish farming has highlighted the need for more efficient management across breeding, feeding, harvesting, inventory, and distribution systems. Digital tools are becoming increasingly important in improving productivity, resource monitoring, and decision-making in aquatic food and farm production systems (Food and Agriculture Organization, 2022; WorldFish, 2021). Aquaculture value chain research also emphasizes the importance of understanding the linkages among producers, input suppliers, market intermediaries, and consumers across production, processing, and distribution activities (Bush et al., 2019).")

    p_body("One of the major challenges faced by superworm and crayfish farms is the lack of centralized systems capable of managing operational data across different production units, species, batches, and growth cycles. Without an integrated digital platform, farm managers may encounter difficulties in consolidating production records, monitoring feed and water usage, tracking mortality and growth performance, and evaluating overall farm productivity. Many small and medium farm enterprises continue to experience inefficient data management and limited technological integration. Digital platforms can support farmers and value chain actors by improving access to production information, feed management, market linkages, and farm performance tracking (WorldFish, 2021; WorldFish, 2022).")

    p_body("To address these challenges, the development of a Superworm and Crayfish Value Chain Management System with Predictive Analytics is proposed to support more efficient data management and value chain coordination. The proposed system aims to provide a centralized digital platform capable of storing, organizing, and analyzing data related to superworm rearing, crayfish pond operations, feeding activities, inventory, harvesting, supply transactions, and market distribution. By incorporating predictive analytics, the system can generate data-driven insights that support planning and strategic decision-making, such as estimating production trends, forecasting harvest volume, monitoring supply levels, and improving resource allocation. Recent research in aquatic production forecasting demonstrates that predictive models using historical and environmental data can support producers and decision-makers in improving planning and resource management across aquaculture-related value chains (Hu et al., 2025).")

    # 1.1 Background of the Study
    p_heading_3("1.1 Background of the Study")
    p_body("Silay Superworm & Crayfish, a farming enterprise located in Brgy. Lantad, Silay City, operates production areas for superworm rearing and crayfish cultivation. As the demand for alternative protein sources and freshwater aquaculture products continues to grow, the farm needs to modernize its operations to meet increasing market demand and manage growing production complexity. Stakeholders, including the farm owner, delivery riders, and customers, require timely and accurate information to monitor production, manage resources, and coordinate activities across different production and distribution channels.")
    p_body("Currently, Silay Superworm & Crayfish faces challenges in managing operational data and coordinating value chain activities due to fragmented record-keeping practices. These include manual logs and paper-based systems, which limit efficient monitoring of superworm batches, crayfish pond performance, inventory tracking, feeding schedules, harvesting activities, and sales records. This fragmented approach reduces transparency and slows down decision-making, potentially leading to inefficient resource allocation and reduced competitiveness.")
    p_body("To address these issues, a centralized digital management system with predictive analytics is proposed. The system will integrate superworm production management, crayfish pond monitoring, feed and harvest data, inventory records, and sales information into a unified web-based platform. By providing real-time insights and predictive forecasts, the system aims to improve coordination, enhance productivity, support better decision-making, and contribute to the sustainable growth of Silay Superworm & Crayfish.")

    # 1.2 Objectives of the Study
    p_heading_3("1.2 Objectives of the Study")
    p_body("The purpose of the study is to develop a Silay Superworm & Crayfish Value Chain Management System with Predictive Analytics that improves the monitoring, management, and analysis of farm operations, including production, inventory, sales, and expenses.")
    p_body("Specifically, this study aims to:")
    
    p_objective_item("1. ", "Design and develop a Silay Superworm & Crayfish Value Chain Management System with the following features:")
    p_objective_item("1.1. ", "Documentation of superworm and crayfish production processes, including breeding, feeding, harvesting, processing, distribution, and delivery", sub=True)
    p_objective_item("1.2. ", "Sales and inventory management for tracking stock, transactions, and product availability", sub=True)
    p_objective_item("1.3. ", "Product classification based on type, category, and scientific name", sub=True)
    p_objective_item("1.4. ", "Monitoring of monthly operational expenses", sub=True)
    p_objective_item("1.5. ", "Forecasting and predictive analytics for production, expenses, and market demand", sub=True)
    p_objective_item("1.6. ", "Customer transaction management, including wholesale, retail, and reservations", sub=True)
    p_objective_item("1.7. ", "Data clustering for analyzing production, sales patterns, and customer behavior", sub=True)
    p_objective_item("1.8. ", "Report generation and analytics for decision-making", sub=True)
    p_objective_item("2. ", "Test the functionality of the developed system.")
    p_objective_item("3. ", "Evaluate the usability of the system in terms of:")
    p_objective_item("3.1. ", "System usefulness", sub=True)
    p_objective_item("3.2. ", "Quality of information", sub=True)
    p_objective_item("3.3. ", "Interface quality", sub=True)
    p_objective_item("3.4. ", "Overall usability", sub=True)
    p_objective_item("4. ", "Develop a User Manual for the system.")

    # 1.3 Scope and Limitations of the Study
    p_heading_3("1.3 Scope and Limitations of the Study")
    p_body("This study focuses on the development of a Silay Superworm & Crayfish Value Chain Management System with Predictive Analytics designed for Silay Superworm & Crayfish. The system provides a centralized platform for managing farm production data, including superworm rearing, crayfish pond operations, harvest records, product classification, inventory, sales transactions, and operational expenses. It also includes predictive analytics to generate forecasts related to production, expenses, and market demand. The system supports customer transaction management and generates analytical reports for decision-making. The system will be web-based and accessible through standard browsers, with data stored in a centralized database.")
    p_body("The study has limitations. The system will not include real-time sensor integration, such as temperature, dissolved oxygen, pH levels, or humidity monitoring, since data will be manually entered by authorized users. Predictive analytics will rely on historical data and will not include advanced machine learning models. Furthermore, the system is tailored specifically for Silay Superworm & Crayfish and may require modifications if applied to other farming or aquaculture enterprises.")

    # 1.4 Significance of the Study
    p_heading_3("1.4 Significance of the Study")
    p_body("The development of the Silay Superworm & Crayfish Value Chain Management System with Predictive Analytics is expected to benefit several stakeholders involved in superworm and crayfish production and distribution operations.")
    p_term_item("Farm Owners. ", "The system will provide a centralized platform for monitoring production, inventory, sales, operational expenses, and overall farm performance. It will help the owner make data-driven decisions by providing organized data, predictive forecasts, and analytical reports.")
    p_term_item("Delivery Riders. ", "The system will provide a dedicated mobile delivery interface to receive order dispatch notifications, view delivery details, navigate customer locations, and record proof of delivery upon fulfillment, ensuring smooth and accountable logistics.")
    p_term_item("Customers and Buyers. ", "By providing an online marketplace with real-time product availability, transparent pricing, convenient ordering, location pinning, and digital payment options, the system makes purchasing superworms and crayfish more accessible, reliable, and convenient.")
    p_term_item("Researchers and Future Developers. ", "This study may serve as a reference for future research related to farm management systems, predictive analytics in agriculture and aquaculture, and digital transformation in small-scale production enterprises.")
    p_term_item("The University and Academic Community. ", "The study contributes to the academic field of information systems by demonstrating how digital technologies and predictive analytics can be applied to improve farm management, production monitoring, and value chain coordination.")

    # 1.5 Definition of Terms
    p_heading_3("1.5 Definition of Terms")
    p_term_item("Aquaculture. ", "The farming of fish, shellfish, and other aquatic organisms in controlled environments such as ponds, tanks, or coastal areas.")
    p_term_item("Aquaculture Management System. ", "A digital platform designed to monitor, manage, and analyze aquaculture operations, including production, inventory, and financial records.")
    p_term_item("Clustering. ", "A data analysis method used to group similar data points or records based on shared characteristics.")
    p_term_item("Inventory Management. ", "The process of monitoring and controlling the storage, quantity, and movement of products within a system.")
    p_term_item("Operational Expenses. ", "Costs incurred in the daily operations of aquaculture farms, such as feed, labor, maintenance, and utilities.")
    p_term_item("Predictive Analytics. ", "A data analysis technique that uses historical data to identify patterns and generate predictions about future outcomes.")
    p_term_item("Retail Sales. ", "The selling of products in smaller quantities directly to individual consumers.")
    p_term_item("Value Chain. ", "A series of activities involved in producing and delivering a product, including production, processing, distribution, and sales.")
    p_term_item("Wholesale Sales. ", "The selling of products in large quantities to retailers, distributors, or businesses.")

    output_path = r'c:\Users\kymer\Desktop\Lantad ACarwfish Django\Documents\Final_Paper\Chapter_1.docx'
    doc.save(output_path)
    print("Successfully built Short Bond Paper (Letter: 8.5 x 11 inches) Chapter_1.docx!")

if __name__ == '__main__':
    create_chapter_1()

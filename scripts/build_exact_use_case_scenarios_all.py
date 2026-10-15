import docx
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import os

# Complete, precise Use Case Scenarios matching the exact format shown in the user images:
ALL_USE_CASES = [
    {
        "table_num": "3",
        "table_title": "Login",
        "uc_title": "Use Case 1: Login",
        "goal": "To allow authorized users (Farm Owner, Customer, and Delivery Rider) to securely authenticate and access their role-specific interfaces.",
        "scenario": "Scenario 1: Successful user authentication.",
        "actor": "Farm Owner, Customer, Delivery Rider",
        "precondition": "The user is registered in the system with an active account and assigned role.",
        "steps": [
            "Step 1: The user opens the system login page.",
            "Step 2: The system displays the authentication form with username and password input fields.",
            "Step 3: The user enters registered credentials and clicks Sign In.",
            "Step 4: The system validates the credentials against the database.",
            "Step 5: The system verifies the user role and generates an authenticated session.",
            "Step 6: The system redirects the user to their role-specific dashboard.",
            "Step 7: The system displays a welcome notification on the user interface."
        ],
        "user_action": "Enters registered credentials and clicks \"Sign In.\"",
        "system_response": "Validates credentials, establishes an authenticated session, and redirects the user to their designated dashboard."
    },
    {
        "table_num": "4",
        "table_title": "View Dashboard",
        "uc_title": "Use Case 2: View Dashboard",
        "goal": "To allow the Farm Owner to view real-time aquaculture operational summaries, pond maintenance alerts, and pending orders.",
        "scenario": "Scenario 1: Successful dashboard overview.",
        "actor": "Farm Owner",
        "precondition": "The user is logged in with executive privileges as Farm Owner.",
        "steps": [
            "Step 1: The user navigates to the Executive Dashboard.",
            "Step 2: The system retrieves active pond statistics, sanitization alerts, feed consumption, and order queues.",
            "Step 3: The system renders summary KPI cards and real-time operational charts.",
            "Step 4: The user reviews active alerts and pending dispatch requests.",
            "Step 5: The user clicks on an operational KPI card to inspect detailed module logs.",
            "Step 6: The system loads the corresponding operational details.",
            "Step 7: The system displays filtered records for monitoring and decision-making."
        ],
        "user_action": "Navigates to the Executive Dashboard and selects operational indicators.",
        "system_response": "Queries real-time aquaculture metrics, renders visual KPI summary cards, and displays detailed operational records."
    },
    {
        "table_num": "5",
        "table_title": "Create Stock Batch",
        "uc_title": "Use Case 3: Create Stock Batch",
        "goal": "To allow the Farm Owner to initialize and register a new aquaculture rearing batch (superworms or redclaw crayfish).",
        "scenario": "Scenario 1: Successful stock batch creation.",
        "actor": "Farm Owner",
        "precondition": "The user is logged in as Farm Owner, and rearing capacity is available.",
        "steps": [
            "Step 1: The user navigates to the Cultivation Batches module.",
            "Step 2: The system displays existing batch rosters and available batch creation actions.",
            "Step 3: The user clicks \"New Batch\" and enters species type, initial stocking quantity, growth stage, and start date.",
            "Step 4: The system validates input completeness and positive numerical quantities.",
            "Step 5: The user confirms the batch creation.",
            "Step 6: The system generates a unique batch identifier and stores the record in the database.",
            "Step 7: The system displays the new batch record in the active batch inventory list."
        ],
        "user_action": "Enters species specifications, stocking count, and confirms batch creation.",
        "system_response": "Validates input values, creates the new Cultivation_Batch record, and updates the active batch list."
    },
    {
        "table_num": "6",
        "table_title": "Assign Ponds",
        "uc_title": "Use Case 4: Assign Ponds",
        "goal": "To allow the Farm Owner to allocate juvenile batches to specific rearing ponds or cultivation shelves.",
        "scenario": "Scenario 1: Successful pond allocation.",
        "actor": "Farm Owner",
        "precondition": "The user is logged in as Farm Owner, and the target pond is sanitized and active.",
        "steps": [
            "Step 1: The user navigates to the Pond Operations module.",
            "Step 2: The system displays physical ponds, current occupancy status, and capacity thresholds.",
            "Step 3: The user selects a pond and assigns an active stock batch.",
            "Step 4: The system checks density limits and pond compatibility.",
            "Step 5: The user confirms the pond assignment.",
            "Step 6: The system updates the pond status to Occupied and links the batch record.",
            "Step 7: The system displays the updated pond status and logs the assignment."
        ],
        "user_action": "Selects target pond, assigns batch parameters, and confirms allocation.",
        "system_response": "Validates pond capacity limits, links the batch to the pond, updates occupancy status, and logs the assignment."
    },
    {
        "table_num": "7",
        "table_title": "Record Feed Administration and Pond Maintenance Logs",
        "uc_title": "Use Case 5: Record Feed Administration and Pond Maintenance Logs",
        "goal": "To allow the Farm Owner to log daily nutritional feed rations and document pond sanitization and maintenance activities across active rearing units.",
        "scenario": "Scenario 1: Successful feed administration and pond maintenance recording.",
        "actor": "Farm Owner",
        "precondition": "The user is logged in as Farm Owner, and active occupied ponds exist.",
        "steps": [
            "Step 1: The user opens the Pond Daily Logs module.",
            "Step 2: The system displays occupied ponds and daily maintenance log options.",
            "Step 3: The user selects a pond and enters feed type, feed amount (kg), and pond maintenance status.",
            "Step 4: The system validates entered numerical values and deducts feed stock from inventory.",
            "Step 5: The user clicks \"Save Daily Log\".",
            "Step 6: The system records entries in Feeding_Log and Operations_Pond maintenance records.",
            "Step 7: The system displays an update confirmation and updates pond maintenance schedules."
        ],
        "user_action": "Enters feed quantities and maintenance details, then clicks \"Save Daily Log.\"",
        "system_response": "Validates log inputs, commits feeding and maintenance records, deducts feed inventory, and updates pond operational status."
    },
    {
        "table_num": "8",
        "table_title": "Manage Inventory and Pricing",
        "uc_title": "Use Case 6: Manage Inventory and Pricing",
        "goal": "To allow the Farm Owner to monitor real-time stock volumes, configure dual pricing (per-piece and per-kg), and adjust reorder levels.",
        "scenario": "Scenario 1: Successful inventory and pricing adjustment.",
        "actor": "Farm Owner",
        "precondition": "The user is logged in as Farm Owner with inventory management privileges.",
        "steps": [
            "Step 1: The user navigates to the Inventory Management module.",
            "Step 2: The system displays commercial products, active stock volumes, and dual pricing tiers.",
            "Step 3: The user selects a product and updates stock quantities or per-piece/per-kg prices.",
            "Step 4: The system validates entered numerical prices and minimum stock thresholds.",
            "Step 5: The user confirms the inventory updates.",
            "Step 6: The system commits modifications to Inventory_Item and generates a Stock_Movement record.",
            "Step 7: The system displays the updated inventory roster and refreshes catalog prices."
        ],
        "user_action": "Modifies product stock or pricing tiers, then confirms updates.",
        "system_response": "Validates pricing inputs, updates product inventory records, creates a Stock_Movement log, and refreshes storefront pricing."
    },
    {
        "table_num": "9",
        "table_title": "Manage Operational Expenses",
        "uc_title": "Use Case 7: Manage Operational Expenses",
        "goal": "To allow the Farm Owner to log, categorize, and track operational expenditures such as feeds, utilities, fuel, and facility maintenance.",
        "scenario": "Scenario 1: Successful operational expense recording.",
        "actor": "Farm Owner",
        "precondition": "The user is logged in as Farm Owner with financial accounting access.",
        "steps": [
            "Step 1: The user navigates to the Expense Management module.",
            "Step 2: The system displays existing expense records and summary expense categories.",
            "Step 3: The user clicks \"Add Expense\" and selects category, enters amount, date, and description.",
            "Step 4: The system validates amount format and required descriptions.",
            "Step 5: The user confirms the expense entry.",
            "Step 6: The system saves the Expense record and updates the expense ledger.",
            "Step 7: The system displays the new expense entry in the financial log list."
        ],
        "user_action": "Enters expense category, amount, and confirms the expense entry.",
        "system_response": "Validates submitted financial values, commits the Expense record, and updates financial ledger summaries."
    },
    {
        "table_num": "10",
        "table_title": "Customer Purchase",
        "uc_title": "Use Case 8: Customer Purchase",
        "goal": "To allow Commercial Customers to browse commercial aquaculture products, configure cart line items, and submit orders with delivery options.",
        "scenario": "Scenario 1: Successful catalog browsing and order placement.",
        "actor": "Customer",
        "precondition": "The customer is on the public storefront, and commercial inventory is active.",
        "steps": [
            "Step 1: The customer browses the product catalog (live crayfish, packaged superworms, feeds).",
            "Step 2: The system displays available stock volumes and dual pricing options (per-piece and per-kg).",
            "Step 3: The customer selects quantity, unit pricing tier, and adds items to the cart.",
            "Step 4: The system calculates subtotals and updates shopping cart state.",
            "Step 5: The customer selects fulfillment method (Delivery / Storefront Pickup) and enters delivery address.",
            "Step 6: The customer selects payment method (GCash / COD) and submits the checkout.",
            "Step 7: The system creates an Order record with status Pending and displays the order summary."
        ],
        "user_action": "Selects products, configures delivery address, and submits checkout.",
        "system_response": "Calculates order total, creates an Order record with status Pending, reserves temporary inventory, and returns order summary."
    },
    {
        "table_num": "11",
        "table_title": "Process Sales Order",
        "uc_title": "Use Case 9: Process Sales Order",
        "goal": "To allow the Farm Owner to review customer order requests, verify GCash payment receipts, and approve orders for dispatch.",
        "scenario": "Scenario 1: Successful sales order approval.",
        "actor": "Farm Owner",
        "precondition": "A customer has placed an order with payment details (GCash receipt upload or COD).",
        "steps": [
            "Step 1: The user navigates to the Orders Processing Queue.",
            "Step 2: The system displays pending customer orders and uploaded payment receipts.",
            "Step 3: The user inspects ordered items, delivery destination, and payment validity.",
            "Step 4: The user clicks \"Approve Order\".",
            "Step 5: The system updates Order status to Confirmed and Payment status to Approved.",
            "Step 6: The system commits permanent stock deduction from inventory.",
            "Step 7: The system queues the confirmed order into the courier dispatch pool."
        ],
        "user_action": "Inspects order details and payment proof, then clicks \"Approve Order.\"",
        "system_response": "Validates payment verification, sets Order status to Confirmed, deducts inventory, and queues order for dispatch."
    },
    {
        "table_num": "12",
        "table_title": "Courier Dispatch",
        "uc_title": "Use Case 10: Courier Dispatch",
        "goal": "To allow the Farm Owner to assign confirmed orders to active Delivery Riders and provide riders with destination waypoints.",
        "scenario": "Scenario 1: Successful courier assignment.",
        "actor": "Farm Owner, Delivery Rider",
        "precondition": "The order is marked Confirmed, and active delivery couriers are available.",
        "steps": [
            "Step 1: The Farm Owner opens the Dispatch Pool and selects an unassigned confirmed order.",
            "Step 2: The system displays active riders on duty.",
            "Step 3: The Farm Owner selects a rider and clicks \"Assign Delivery\".",
            "Step 4: The system creates a Delivery_Dispatch record and generates route waypoints.",
            "Step 5: The system alerts the assigned Delivery Rider with customer drop-off details.",
            "Step 6: The Delivery Rider views the assigned delivery order on their mobile interface.",
            "Step 7: The system displays the dispatch confirmation and updates order status to Dispatched."
        ],
        "user_action": "Selects available courier and clicks \"Assign Delivery.\"",
        "system_response": "Creates dispatch assignment, updates delivery status to Dispatched, routes waypoints, and alerts the assigned courier."
    },
    {
        "table_num": "13",
        "table_title": "Confirm Delivery (Proof of Delivery)",
        "uc_title": "Use Case 11: Confirm Delivery (Proof of Delivery)",
        "goal": "To allow the Delivery Rider to capture recipient proof (POD photo/signature), collect COD payments, and finalize delivery.",
        "scenario": "Scenario 1: Successful delivery completion with POD.",
        "actor": "Delivery Rider",
        "precondition": "The order is in transit, and the rider has reached the customer drop-off location.",
        "steps": [
            "Step 1: The Delivery Rider hands over the package and collects COD cash if applicable.",
            "Step 2: The Delivery Rider opens the active order on their mobile interface.",
            "Step 3: The Delivery Rider captures Proof of Delivery (recipient photo or signature).",
            "Step 4: The Delivery Rider clicks \"Complete Delivery\".",
            "Step 5: The system uploads the POD evidence and updates Order status to Delivered.",
            "Step 6: The system records the completion timestamp in the Delivery_Dispatch log.",
            "Step 7: The system notifies the Customer and Farm Owner of successful fulfillment."
        ],
        "user_action": "Uploads POD photo or signature and clicks \"Complete Delivery.\"",
        "system_response": "Validates POD submission, updates Order status to Delivered, records completion timestamp, and sends delivery alerts."
    },
    {
        "table_num": "14",
        "table_title": "Export Report",
        "uc_title": "Use Case 12: Export Report",
        "goal": "To allow the Farm Owner to compile, preview, and export official printable reports for sales, expenses, and inventory balances.",
        "scenario": "Scenario 1: Successful report generation and export.",
        "actor": "Farm Owner",
        "precondition": "The user is logged in as Farm Owner, and operational records exist.",
        "steps": [
            "Step 1: The user navigates to the Reports & Analytics module.",
            "Step 2: The user selects report category (Sales Summary, Expense Breakdown, Inventory Log).",
            "Step 3: The user defines date range filters (daily, weekly, monthly) and clicks \"Generate Report\".",
            "Step 4: The system queries database tables and compiles aggregate statistics.",
            "Step 5: The system renders the formatted report preview.",
            "Step 6: The user reviews the preview and clicks \"Export PDF / Print\".",
            "Step 7: The system generates the formatted document and initiates file download."
        ],
        "user_action": "Selects report category, applies date range, and clicks \"Export PDF / Print.\"",
        "system_response": "Compiles transaction datasets, renders printable report preview, and generates the downloadable export file."
    },
    {
        "table_num": "15",
        "table_title": "Sales Forecasting and Clustering",
        "uc_title": "Use Case 13: Sales Forecasting and Clustering",
        "goal": "To allow the Farm Owner to execute predictive machine learning algorithms for short-term sales revenue forecasting and customer segmentation.",
        "scenario": "Scenario 1: Successful predictive analytics computation.",
        "actor": "Farm Owner",
        "precondition": "Historical sales and harvesting records are present in the database.",
        "steps": [
            "Step 1: The user navigates to the Predictive Analytics module.",
            "Step 2: The system loads historical sales records, order volumes, and harvest yields.",
            "Step 3: The user clicks \"Run Predictive Models\".",
            "Step 4: The system executes Linear Regression algorithms to forecast sales trends.",
            "Step 5: The system runs K-Means clustering to segment customer purchasing cohorts.",
            "Step 6: The system visualizes projected revenue graphs and regional demand heatmaps.",
            "Step 7: The system displays generated forecasts and clustering insights for strategic planning."
        ],
        "user_action": "Navigates to Predictive Analytics and clicks \"Run Predictive Models.\"",
        "system_response": "Processes historical datasets through regression and clustering models, rendering forecast visualizations and segmentation reports."
    }
]

def render_exact_use_case_table(doc, uc):
    # Table Title (e.g. Table 16 Generate Yearly Record Archive)
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(3)
    p_title.paragraph_format.line_spacing = 1.0
    r_t = p_title.add_run(f"Table {uc['table_num']} {uc['table_title']}")
    r_t.font.name = "Arial"
    r_t.font.size = Pt(11)
    r_t.bold = True

    # 1-Column Table (2 rows: Row 0 is Use Case Title, Row 1 is Content)
    table = doc.add_table(rows=2, cols=1)
    table.autofit = False
    table.columns[0].width = Inches(6.0)

    # Set exact APA single borders (top, middle horizontal, bottom; no left/right/vertical)
    tblPr = table._tbl.tblPr
    tblBorders = parse_xml(
        f'<w:tblBorders {nsdecls("w")}>\n'
        f'<w:top w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:left w:val="nil"/>\n'
        f'<w:bottom w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:right w:val="nil"/>\n'
        f'<w:insideH w:val="single" w:sz="8" w:space="0" w:color="000000"/>\n'
        f'<w:insideV w:val="nil"/>\n'
        f'</w:tblBorders>'
    )
    tblPr.append(tblBorders)

    # Row 0: Use Case Header (e.g. Use Case 14: Generate Yearly Record Archive)
    cell_0 = table.rows[0].cells[0]
    p0 = cell_0.paragraphs[0]
    p0.paragraph_format.space_before = Pt(3)
    p0.paragraph_format.space_after = Pt(3)
    p0.paragraph_format.line_spacing = 1.0
    r0 = p0.add_run(uc['uc_title'])
    r0.font.name = "Arial"
    r0.font.size = Pt(10)
    r0.bold = True

    # Row 1: Content
    cell_1 = table.rows[1].cells[0]

    # Goal
    p_goal = cell_1.paragraphs[0]
    p_goal.paragraph_format.space_before = Pt(3)
    p_goal.paragraph_format.space_after = Pt(1)
    p_goal.paragraph_format.line_spacing = 1.0
    rg1 = p_goal.add_run("Use Case Goal: ")
    rg1.bold = True
    rg1.font.name = "Arial"
    rg1.font.size = Pt(9.5)
    rg2 = p_goal.add_run(uc['goal'])
    rg2.font.name = "Arial"
    rg2.font.size = Pt(9.5)

    # Scenario Title
    p_sc = cell_1.add_paragraph()
    p_sc.paragraph_format.space_before = Pt(1)
    p_sc.paragraph_format.space_after = Pt(1)
    p_sc.paragraph_format.line_spacing = 1.0
    r_sc = p_sc.add_run(uc['scenario'])
    r_sc.bold = True
    r_sc.font.name = "Arial"
    r_sc.font.size = Pt(9.5)

    # Actor
    p_act = cell_1.add_paragraph()
    p_act.paragraph_format.space_before = Pt(0)
    p_act.paragraph_format.space_after = Pt(1)
    p_act.paragraph_format.line_spacing = 1.0
    r_lbl = p_act.add_run("Actor: ")
    r_lbl.bold = True
    r_lbl.font.name = "Arial"
    r_lbl.font.size = Pt(9.5)
    r_val = p_act.add_run(uc['actor'])
    r_val.font.name = "Arial"
    r_val.font.size = Pt(9.5)

    # Pre-condition
    p_pre = cell_1.add_paragraph()
    p_pre.paragraph_format.space_before = Pt(0)
    p_pre.paragraph_format.space_after = Pt(1)
    p_pre.paragraph_format.line_spacing = 1.0
    r_lbl = p_pre.add_run("Pre-condition: ")
    r_lbl.bold = True
    r_lbl.font.name = "Arial"
    r_lbl.font.size = Pt(9.5)
    r_val = p_pre.add_run(uc['precondition'])
    r_val.font.name = "Arial"
    r_val.font.size = Pt(9.5)

    # Main Flow
    p_mf = cell_1.add_paragraph()
    p_mf.paragraph_format.space_before = Pt(0)
    p_mf.paragraph_format.space_after = Pt(1)
    p_mf.paragraph_format.line_spacing = 1.0
    r_lbl = p_mf.add_run("Main Flow:")
    r_lbl.bold = True
    r_lbl.font.name = "Arial"
    r_lbl.font.size = Pt(9.5)

    # Steps
    for step in uc['steps']:
        p_st = cell_1.add_paragraph()
        p_st.paragraph_format.space_before = Pt(0)
        p_st.paragraph_format.space_after = Pt(1)
        p_st.paragraph_format.line_spacing = 1.0
        parts = step.split(':', 1)
        if len(parts) == 2:
            r_lbl = p_st.add_run(parts[0] + ':')
            r_lbl.bold = True
            r_lbl.font.name = "Arial"
            r_lbl.font.size = Pt(9.5)
            r_val = p_st.add_run(parts[1])
            r_val.font.name = "Arial"
            r_val.font.size = Pt(9.5)
        else:
            r_st = p_st.add_run(step)
            r_st.font.name = "Arial"
            r_st.font.size = Pt(9.5)

    # User Action
    p_ua = cell_1.add_paragraph()
    p_ua.paragraph_format.space_before = Pt(0)
    p_ua.paragraph_format.space_after = Pt(1)
    p_ua.paragraph_format.line_spacing = 1.0
    r_ua_lbl = p_ua.add_run("User Action: ")
    r_ua_lbl.bold = True
    r_ua_lbl.font.name = "Arial"
    r_ua_lbl.font.size = Pt(9.5)
    r_ua_val = p_ua.add_run(uc['user_action'])
    r_ua_val.font.name = "Arial"
    r_ua_val.font.size = Pt(9.5)

    # System Response
    p_sr = cell_1.add_paragraph()
    p_sr.paragraph_format.space_before = Pt(0)
    p_sr.paragraph_format.space_after = Pt(3)
    p_sr.paragraph_format.line_spacing = 1.0
    r_sr_lbl = p_sr.add_run("System Response: ")
    r_sr_lbl.bold = True
    r_sr_lbl.font.name = "Arial"
    r_sr_lbl.font.size = Pt(9.5)
    r_sr_val = p_sr.add_run(uc['system_response'])
    r_sr_val.font.name = "Arial"
    r_sr_val.font.size = Pt(9.5)

    # Spacing after table
    p_sp = doc.add_paragraph()
    p_sp.paragraph_format.space_before = Pt(0)
    p_sp.paragraph_format.space_after = Pt(6)
    p_sp.paragraph_format.line_spacing = 1.0

def generate_use_case_docs():
    # 1. Generate standalone Documents/Use_Case_Scenario.docx
    doc_uc = docx.Document()
    for section in doc_uc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)

    p_h = doc_uc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(0)
    p_h.paragraph_format.space_after = Pt(6)
    p_h.paragraph_format.line_spacing = 1.5
    r_h = p_h.add_run("3.6 Use Case Scenario")
    r_h.font.name = "Arial"
    r_h.font.size = Pt(12)
    r_h.bold = True

    p_i = doc_uc.add_paragraph()
    p_i.paragraph_format.space_before = Pt(0)
    p_i.paragraph_format.space_after = Pt(6)
    p_i.paragraph_format.line_spacing = 1.5
    p_i.paragraph_format.first_line_indent = Inches(0.5)
    r_i = p_i.add_run(
        "The Use Case Scenario describes the interaction between the primary entities (Farm Owner, "
        "Delivery Rider, and Commercial Customers) and the Aquaculture Value Chain Management System with Predictive Analytics. "
        "Each table details the operational goal, prerequisites, step-by-step main flow of activities, "
        "and summary user action and system response across farm operations, inventory management, order processing, "
        "courier dispatch, and predictive analytics forecasting."
    )
    r_i.font.name = "Arial"
    r_i.font.size = Pt(12)

    for uc in ALL_USE_CASES:
        render_exact_use_case_table(doc_uc, uc)

    doc_uc.save("Documents/Use_Case_Scenario.docx")
    print(f"Successfully saved Documents/Use_Case_Scenario.docx with {len(ALL_USE_CASES)} exact Use Case Tables!")

if __name__ == '__main__':
    generate_use_case_docs()

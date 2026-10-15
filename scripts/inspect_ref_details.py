import docx

doc = docx.Document(r'c:\Users\kymer\Desktop\Lantad ACarwfish Django\Documents\Final_Paper\Herbias, Hulleza, Opjer - GovSupply Connect_ A Municipal Procurement Transparency and Supply Distribution Monitoring System.docx')

for i in range(len(doc.paragraphs)):
    p = doc.paragraphs[i]
    if any(k in p.text for k in ['CHAPTER', 'Introduction', 'Background', 'Objectives', 'Significance', 'Scope', 'Definition', 'GOVSUPPLY']):
        print(f"P{i:03d} [Style={p.style.name}, Align={p.alignment}]: '{p.text}'")
        print(f"     Format: space_before={p.paragraph_format.space_before}, space_after={p.paragraph_format.space_after}, line_spacing={p.paragraph_format.line_spacing}, indent={p.paragraph_format.first_line_indent}")
        for r in p.runs:
            print(f"     Run: text='{r.text}', font={r.font.name}, size={r.font.size.pt if r.font.size else None}, bold={r.bold}, italic={r.italic}")

import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def add_section_3_8(doc):
    # Heading 3.8 Hardware and Software Requirements
    p_h = doc.add_paragraph()
    p_h.paragraph_format.space_before = Pt(14)
    p_h.paragraph_format.space_after = Pt(6)
    p_h.paragraph_format.line_spacing = 1.5
    r_h = p_h.add_run("3.8 Hardware and Software Requirements")
    r_h.font.name = "Arial"
    r_h.font.size = Pt(12)
    r_h.font.bold = True
    r_h.font.color.rgb = RGBColor(0, 0, 0)

    # Intro Paragraph
    p_intro = doc.add_paragraph()
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(12)
    p_intro.paragraph_format.line_spacing = 1.5
    p_intro.paragraph_format.first_line_indent = Inches(0.5)
    r_intro = p_intro.add_run(
        "Hardware and Software Requirements provide the detailed minimum requirements needed "
        "to support the functions and operation of the developed system."
    )
    r_intro.font.name = "Arial"
    r_intro.font.size = Pt(12)
    r_intro.font.color.rgb = RGBColor(0, 0, 0)

    # Subheading: Recommended Minimum Hardware Requirements
    p_sub_hw = doc.add_paragraph()
    p_sub_hw.paragraph_format.space_before = Pt(8)
    p_sub_hw.paragraph_format.space_after = Pt(6)
    p_sub_hw.paragraph_format.line_spacing = 1.5
    r_sub_hw = p_sub_hw.add_run("Recommended Minimum Hardware Requirements")
    r_sub_hw.font.name = "Arial"
    r_sub_hw.font.size = Pt(12)
    r_sub_hw.font.bold = True
    r_sub_hw.font.color.rgb = RGBColor(0, 0, 0)

    hw_items = [
        ("Computer: ", "Desktop or laptop computer"),
        ("Processor: ", "Intel Core i3 or higher (or AMD Ryzen 3 equivalent)"),
        ("Processor Speed: ", "3.00 GHz or above"),
        ("RAM: ", "At least 8 GB of RAM for smooth system performance"),
        ("Storage: ", "At least 256 GB HDD or SSD"),
        ("Monitor: ", "Standard monitor with at least 1366 × 768 screen resolution"),
        ("Mobile Device (for Delivery Rider & Mobile Users): ", "Smartphone running Android 10.0+ or iOS 14.0+ with camera (Proof of Delivery) and GPS capability"),
        ("Internet Connection: ", "Stable internet connection for accessing the web-based system"),
    ]

    for label, desc in hw_items:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_before = Pt(0)
        p_item.paragraph_format.space_after = Pt(4)
        p_item.paragraph_format.line_spacing = 1.5
        p_item.paragraph_format.left_indent = Inches(0.75)
        
        r_lbl = p_item.add_run(label)
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(12)
        r_lbl.font.bold = False
        r_lbl.font.color.rgb = RGBColor(0, 0, 0)

        r_desc = p_item.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(12)
        r_desc.font.color.rgb = RGBColor(0, 0, 0)

    # Subheading: Recommended Software Requirements
    p_sub_sw = doc.add_paragraph()
    p_sub_sw.paragraph_format.space_before = Pt(12)
    p_sub_sw.paragraph_format.space_after = Pt(6)
    p_sub_sw.paragraph_format.line_spacing = 1.5
    r_sub_sw = p_sub_sw.add_run("Recommended Software Requirements")
    r_sub_sw.font.name = "Arial"
    r_sub_sw.font.size = Pt(12)
    r_sub_sw.font.bold = True
    r_sub_sw.font.color.rgb = RGBColor(0, 0, 0)

    sw_items = [
        ("Operating System: ", "Windows 10, Windows 11, or Linux (Ubuntu Server 22.04 LTS / 24.04 LTS)"),
        ("Web Browser: ", "Google Chrome, Microsoft Edge, Mozilla Firefox, or any modern web browser"),
        ("Programming Language & Backend: ", "Python 3.11 with Django 5.1 Web Framework"),
        ("Database Management System: ", "PostgreSQL 16 / SQLite 3 relational database"),
        ("Predictive Analytics & Machine Learning Libraries: ", "Scikit-Learn, NumPy, Pandas, and SciPy (Linear Regression and K-Means Clustering)"),
        ("Frontend Technologies: ", "HTML5, CSS3, JavaScript (ES6+), and Bootstrap 5 / Tailwind CSS"),
        ("Interactive Visualization Libraries: ", "Chart.js, ApexCharts, and Leaflet.js (OpenStreetMap GIS Mapping)"),
        ("Web Server & WSGI Gateway: ", "Gunicorn WSGI server with Nginx HTTP reverse proxy"),
    ]

    for label, desc in sw_items:
        p_item = doc.add_paragraph(style='List Bullet')
        p_item.paragraph_format.space_before = Pt(0)
        p_item.paragraph_format.space_after = Pt(4)
        p_item.paragraph_format.line_spacing = 1.5
        p_item.paragraph_format.left_indent = Inches(0.75)
        
        r_lbl = p_item.add_run(label)
        r_lbl.font.name = "Arial"
        r_lbl.font.size = Pt(12)
        r_lbl.font.bold = False
        r_lbl.font.color.rgb = RGBColor(0, 0, 0)

        r_desc = p_item.add_run(desc)
        r_desc.font.name = "Arial"
        r_desc.font.size = Pt(12)
        r_desc.font.color.rgb = RGBColor(0, 0, 0)

def build_standalone():
    doc = Document()
    for section in doc.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)
        section.left_margin = Inches(1.5)
        section.right_margin = Inches(1.0)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)

    add_section_3_8(doc)
    out_path = os.path.abspath("Documents/Hardware_and_Software_Requirements_3_8.docx")
    os.makedirs(os.path.dirname(out_path), exist_ok=True)
    doc.save(out_path)
    print(f"Saved: {out_path}")

if __name__ == "__main__":
    build_standalone()
